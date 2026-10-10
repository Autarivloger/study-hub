"""Add Week 24 only; run all examples before the atomic content write."""
import json
import os
from pathlib import Path
import subprocess
import sys
import textwrap
import tempfile

ROOT = Path(__file__).resolve().parent.parent
days = {f"24.{d}": {"parts": []} for d in range(6)}
executed = 0


def clean(code):
    return textwrap.dedent(code).strip() + "\n"


def output(code):
    global executed
    with tempfile.TemporaryDirectory(prefix="studyhub-w24-") as folder:
        run = subprocess.run(
            [sys.executable, "-c", clean(code)], capture_output=True,
            encoding="utf-8", timeout=8, cwd=folder,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
    if run.returncode:
        raise RuntimeError(run.stderr + "\n" + code)
    executed += 1
    return run.stdout.replace("\r\n", "\n").rstrip()


def section(kind, value=None, **extra):
    result = {"t": kind, **extra}
    if value is not None:
        result["v"] = value
    return result


def question(prompt, answer, wrong1, wrong2, why, rotation):
    options = [answer, wrong1, wrong2]
    for _ in range(rotation % 3):
        options.insert(0, options.pop())
    return {"q": prompt, "options": options, "correct": rotation % 3, "why": why}


def add(day, title, goal, concepts, explanation, syntax, code, walkthrough,
        mistake, task, solution, challenge, challenge_code, key, quizzes):
    parts = days[f"24.{day}"]["parts"]
    index = len(parts)
    qs = [question(*q, index + n) for n, q in enumerate(quizzes)]
    sections = [
        section("h", title),
        section("key", "Learning objective: " + goal),
        *[section("new", name, note=note) for name, note in concepts],
        section("p", explanation),
        section("syn", syntax.replace(r"\n", "\n"), note="Read the pattern, then trace the complete example below."),
        section("code", clean(code)),
        section("out", output(code)),
        section("p", walkthrough),
        section("mis", wrong=mistake[0], right=mistake[1], why=mistake[2]),
        section("ex", task, practiceId=f"24.{day}.{index}.build"),
        section("sol", code=clean(solution), out=output(solution),
                why="Build walkthrough: " + key),
        section("ex", "Mini challenge: " + challenge,
                practiceId=f"24.{day}.{index}.challenge"),
        section("sol", code=clean(challenge_code), out=output(challenge_code),
                why="Challenge check: " + key),
        section("try", q=quizzes[0][0], code="# Explain your prediction before revealing.",
                a=quizzes[0][1], why=quizzes[0][4]),
        section("checkpoint", lesson={"id": f"course:24.{day}.{index}", "quiz": qs}),
        section("key", key),
    ]
    parts.append({"title": title, "sections": sections})


add(0, "1. Workbooks, worksheets and cells",
    "Create a small workbook and distinguish its sheets, cell addresses and values.",
    [
        ("openpyxl; .xlsx", "openpyxl reads/writes Excel's .xlsx format without needing Excel installed. It is an installed Python package, not a built-in module; it does not support old binary .xls files."),
        ("Workbook; worksheet; cell", "A workbook is the complete file, a worksheet is one tab, and a cell is a row/column position within that tab."),
        ("Workbook.active; create_sheet; cell coordinates", "active gets the selected worksheet. A1 is column A/row 1; cell(row=1,column=1) uses 1-based coordinates, not Python's zero-based list indexing."),
        ("append; save", "append adds a row of values; save writes the workbook. Saving over an existing filename overwrites it, so use practice copies and new names."),
    ],
    "Week 23 established validated data. Week 24 turns that data into useful documents. Use Python 3.11+ and install "
    "openpyxl, pypdf, reportlab and python-docx once in your practice environment. Reading this page needs none of those packages and works offline. "
    "Start with a tiny workbook in memory; do not edit your only copy of a real office file. A new Workbook already has one sheet. "
    "Library methods are ordinary Python methods like the ones you learned in Weeks 18–21.",
    'from openpyxl import Workbook\nbook = Workbook()\nsheet = book.active\nsheet["A1"] = "Item"\nsheet.append(["Book", 3])',
    r'''
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["Item", "Quantity"])
    sheet.append(["Book", 3])
    summary = book.create_sheet("Summary")
    summary["A1"] = "Sales overview"
    print(book.sheetnames)
    print(sheet["A2"].value)
    print(sheet.cell(row=2, column=2).value)
    ''',
    "Workbook file-जस्तो container हो, Sales एउटा sheet हो। A2 मा Book र B2 मा 3 राखिए। "
    "cell(row=2,column=2) र B2 एउटै स्थान हुन्। Sheets create गर्दा original active sheet मेटिँदैन। "
    "Cell object र त्यसको .value अलग हुन्; print(cell) भन्दा print(cell.value) प्रायः चाहिएको हुन्छ।",
    ("sheet.cell(row=0, column=0)", "sheet.cell(row=1, column=1)",
     "Spreadsheet coordinates start at 1; row/column zero is not A1."),
    "Create Students and Summary sheets. Put ID 007 as text and score 90 as a number; print their types.",
    r'''
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Students"
    sheet.append(["ID", "Score"])
    sheet.append(["007", 90])
    book.create_sheet("Summary")
    print(book.sheetnames)
    print(type(sheet["A2"].value).__name__, type(sheet["B2"].value).__name__)
    ''',
    "Write the same value using A1 notation and row/column notation. Explain why they refer to one cell.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet["C3"] = "Python"
    print(sheet.cell(row=3, column=3).value)
    print(sheet["C3"].coordinate)
    ''',
    "Keep workbook, worksheet, cell and value distinct. Coordinates start at 1. Preserve text IDs. Save into practice copies; openpyxl is not Excel's calculation engine.",
    [
        ("Which coordinates refer to A1?", "row=1, column=1", "row=0, column=0", "row=1, column=0", "Worksheet row/column coordinates are 1-based."),
        ("Does openpyxl require Excel to be installed to create .xlsx?", "No", "Yes", "Only for cells", "It creates the file format directly, although displaying/recalculating it is a separate concern."),
    ])

add(0, "2. Save, load and iterate over rows",
    "Round-trip a workbook and read values without confusing cells with raw values.",
    [
        ("load_workbook", "load_workbook opens a supported existing workbook. Choose its sheet by name rather than assuming the currently active tab is the data tab."),
        ("iter_rows; values_only=True", "iter_rows yields rows of cells; values_only=True yields the cell values instead. Bound the range to the table you need."),
        ("TemporaryDirectory; Path", "The examples use an isolated temporary folder so you can run them without overwriting personal files. Path joins folder/file names."),
        ("read_only=True; close", "Read-only mode can reduce memory for large reads; explicitly close the loaded workbook when finished."),
    ],
    "Create the input inside the example so it works without downloading a hidden sample. Save, load, select Data and iterate from row 2. "
    "values_only=True avoids extra .value calls. Empty cells return None, not zero. max_row/max_column can include formatting-only cells, "
    "so they are not automatically a trustworthy count of valid business records. Read-only workbooks need closing; later parts validate row shape.",
    'book.save(path)\nloaded = load_workbook(path, read_only=True)\nfor row in loaded["Data"].iter_rows(min_row=2, values_only=True):\n    print(row)\nloaded.close()',
    r'''
    import tempfile
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "practice.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet.append(["ID", "Quantity"])
        sheet.append(["007", 3])
        sheet.append(["008", None])
        book.save(path)
        loaded = load_workbook(path, read_only=True)
        try:
            for row in loaded["Data"].iter_rows(min_row=2, max_col=2, values_only=True):
                print(row)
        finally:
            loaded.close()
    ''',
    "Temporary workbook self-contained input हो। Load गरेपछि header skip गर्न min_row=2 छ। "
    "008 को quantity missing भएकाले None आउँछ; यो 0 हो भनेर नमान। finally ले failure भए पनि file resources close गर्छ।",
    ("Use active sheet and assume every styled row is valid data",
     "Select the expected sheet and validate the required table rows",
     "The active tab and worksheet dimensions can reflect viewing/formatting choices rather than your dataset."),
    "Sum quantities from a known valid table of 3 and 5 using iter_rows(values_only=True).",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet.append(["Item", "Quantity"])
    sheet.append(["Book", 3])
    sheet.append(["Pen", 5])
    values = [row[1] for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True)]
    print(values, sum(values))
    ''',
    "Compare normal cell iteration with values_only iteration and print the first cell's coordinate.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet.append(["A", 10])
    cells = next(sheet.iter_rows(max_row=1, max_col=2))
    values = next(sheet.iter_rows(max_row=1, max_col=2, values_only=True))
    print(cells[0].coordinate, cells[0].value)
    print(values)
    ''',
    "Select the expected sheet, bound your read and choose cells versus values intentionally. Missing cells are None. Close read-only workbooks and validate records before aggregation.",
    [
        ("What does values_only=True return?", "Raw cell values in each row", "Only strings", "Only cell coordinates", "It removes the Cell-object wrapper from the iteration result."),
        ("Is max_row always the number of valid data records?", "No", "Yes", "Only if the file has a title", "Formatting and other worksheet contents can extend its dimensions."),
    ])

add(0, "3. Validate spreadsheet rows before computing totals",
    "Reject missing, boolean and malformed quantities rather than silently changing data.",
    [
        ("Typed Excel values", "Cells can contain strings, numbers, booleans, dates, formulas or None. Reading a cell does not validate it against your application schema."),
        ("Actual-integer policy", "type(value) is int rejects booleans; bool is a subclass of int in Python. This example requires nonnegative integer quantities."),
        ("Header contract", "An importer should verify expected column names/order and the data sheet before interpreting rows."),
    ],
    "Our table has exactly Item and Quantity columns. Item must be nonempty text; quantity must be an actual nonnegative integer. "
    "Do not replace every None with 0 or call int on every value: that can hide missing fields and truncate decimals. "
    "For a permissive importer, define accepted conversions explicitly; our first version is strict. "
    "Record rejected worksheet row numbers so the source can be corrected.",
    'check headers\nfor row_number, values in enumerate(rows, 2):\n    check text, exact integer type and range\n    keep rejection row numbers',
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet.append(["Item", "Quantity"])
    for row in [("Book", 3), ("Pen", None), ("Bag", True), ("Cup", -1)]:
        sheet.append(row)
    headers = next(sheet.iter_rows(max_row=1, max_col=2, values_only=True))
    if headers != ("Item", "Quantity"):
        raise ValueError("wrong headers")
    accepted, rejected = [], []
    for number, (item, quantity) in enumerate(
            sheet.iter_rows(min_row=2, max_col=2, values_only=True), 2):
        if not isinstance(item, str) or not item.strip() or type(quantity) is not int or quantity < 0:
            rejected.append(number)
        else:
            accepted.append((item.strip(), quantity))
    print(accepted)
    print("rejected worksheet rows:", rejected)
    print("quantity total:", sum(quantity for item, quantity in accepted))
    ''',
    "Data rows 2 बाट सुरु भएका छन्, त्यसैले rejection list file मा भेट्न सजिलो हुन्छ। "
    "None missing हो, True integer quantity होइन, -1 range बाहिर हो। Book मात्र accepted भएकाले total 3 हुन्छ। "
    "Invalid data हटायौँ भने त्यसको report पनि दिनुपर्छ; silently total घटाउने होइन।",
    ("int(value or 0) accepts every cell safely", "Validate type and domain before aggregation",
     "This shortcut merges missing with zero, accepts booleans and may truncate floats."),
    "Test a checked_quantity function against 0, 3, True, 2.5 and None.",
    r'''
    def checked_quantity(value):
        if type(value) is not int or value < 0:
            raise ValueError("nonnegative integer required")
        return value
    for value in (0, 3, True, 2.5, None):
        try:
            print(repr(value), checked_quantity(value))
        except ValueError:
            print(repr(value), "rejected")
    ''',
    "Validate two different header rows before reading their quantities.",
    r'''
    from openpyxl import Workbook
    for header in (["Item", "Quantity"], ["Quantity", "Item"]):
        sheet = Workbook().active
        sheet.append(header)
        actual = next(sheet.iter_rows(max_row=1, max_col=2, values_only=True))
        print(actual, actual == ("Item", "Quantity"))
    ''',
    "Parsing cell values and validating business records are different jobs. Check headers and types; keep a rejection report. Distinguish actual zero from missing or malformed values.",
    [
        ("Should bool True be accepted as quantity under this policy?", "No", "Yes, because isinstance(True,int)", "Only in Excel", "The declared policy requires an actual int, excluding bool."),
        ("Why retain rejected worksheet row numbers?", "So the original data can be located and corrected.", "To encrypt the file.", "To change headers.", "A total without a clear rejection report can conceal missing source records."),
    ])

add(0, "4. A separate summary sheet and safe output names",
    "Build a summary from validated values and leave the original file unchanged.",
    [
        ("Derived summary", "A summary is computed from validated source rows. It should agree with the same totals used by other reports."),
        ("Source copy vs generated workbook", "openpyxl does not preserve every possible Excel feature when loading and saving. Use separate generated reports rather than overwriting complex originals."),
    ],
    "Produce a Raw sheet and Summary sheet with total quantity and record count. "
    "Always save to a deliberately chosen new filename. Original files may contain macros, shapes or unsupported features; round-tripping through a library can lose them. "
    "For important production output, write to a temporary/staging location, reopen to check contents, then publish. "
    "Our isolated example demonstrates a small supported workbook, not every office file feature.",
    'summary = book.create_sheet("Summary")\nsummary.append(["Metric", "Value"])\nsummary.append(["Total quantity", total])\nbook.save(new_report_path)',
    r'''
    import tempfile
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    rows = [("Book", 3), ("Pen", 5)]
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "generated-report.xlsx"
        book = Workbook()
        raw = book.active
        raw.title = "Raw"
        raw.append(["Item", "Quantity"])
        for row in rows:
            raw.append(row)
        summary = book.create_sheet("Summary")
        summary.append(["Metric", "Value"])
        summary.append(["Total quantity", sum(q for item, q in rows)])
        summary.append(["Records", len(rows)])
        book.save(path)
        restored = load_workbook(path)
        print(restored.sheetnames)
        print(restored["Summary"]["B2"].value, restored["Summary"]["B3"].value)
        restored.close()
    ''',
    "Raw rows बाट Summary values निकालिएका छन्। Saved report फेरि load गर्दा total 8 र count 2 मिल्छ। "
    "नयाँ output को नाम source भन्दा फरक राख; 'save भयो' मात्र होइन 'सही values लेखिए' पनि जाँच।",
    ("Save modified data over the only original workbook", "Write a separate generated report and verify it",
     "Save can overwrite without warning, and unsupported workbook features may not survive."),
    "Compute total and count once from three valid records and use them in a summary dictionary.",
    r'''
    rows = [("Book", 3), ("Pen", 5), ("Bag", 2)]
    metrics = {"total_quantity": sum(q for item, q in rows), "records": len(rows)}
    print(metrics)
    ''',
    "Check that a new output name differs from its source path before saving.",
    r'''
    from pathlib import Path
    source = Path("source.xlsx")
    for output in (Path("report.xlsx"), Path("source.xlsx")):
        print(output.name, source.resolve() != output.resolve())
    ''',
    "Compute shared summary values once, generate separate reports and reopen them for verification. Protect original files and do not promise preservation of every unsupported Excel feature.",
    [
        ("Does Workbook.save automatically protect an existing file?", "No; it can overwrite it.", "Yes", "Only on Windows", "Use a new destination and an explicit overwrite policy."),
        ("Why verify saved summary values?", "Writing a file does not prove it contains the intended totals.", "To install Excel.", "To remove raw data.", "Reopening catches schema, address and calculation mistakes."),
    ])

add(1, "1. Styling: fonts, fills, alignment and number formats",
    "Make a report readable without changing the underlying values.",
    [
        ("Font; PatternFill; Alignment", "openpyxl style objects describe font appearance, cell backgrounds and text placement. Assign them to cells; replace style objects rather than mutating shared styles in place."),
        ("number_format", "A display format controls how a stored value is shown in spreadsheet software. It does not convert or round the stored Python value."),
        ("column_dimensions; row_dimensions", "These configure column widths and row heights. They affect layout, not the data schema."),
    ],
    "Use a clear header, restrained colours, readable column widths and a consistent numeric display. "
    "A format such as 0.00 changes the display of 12.5 to two decimal places in the viewer; the stored value stays 12.5. "
    "Do not use formatting to hide invalid values or disguise the data's units. Assign an 8-digit ARGB colour explicitly. "
    "Styles are not a substitute for validation or arithmetic.",
    'cell.font = Font(bold=True, color="FFFFFFFF")\ncell.fill = PatternFill("solid", fgColor="FF16324F")\ncell.number_format = "0.00"',
    r'''
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    sheet = Workbook().active
    sheet.append(["Item", "Amount"])
    sheet.append(["Book", 12.5])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF16324F")
        cell.alignment = Alignment(horizontal="center")
    sheet["B2"].number_format = "0.00"
    sheet.column_dimensions["A"].width = 24
    print(sheet["A1"].font.bold, sheet["A1"].fill.fgColor.rgb)
    print(sheet["B2"].value, sheet["B2"].number_format)
    ''',
    "Header cells मा एउटै visual rule लगाइएको छ। Amount display 0.00 भए पनि value 12.5 नै हो। "
    "Text पढ्न column width बढायौँ। Report मा amount कुन unit हो भन्ने label छुटाउन हुँदैन।",
    ("number_format='0.00' validates or rounds the actual stored value",
     "Validate and compute values separately; number_format changes display",
     "Formatting is a viewer instruction, not a numeric conversion."),
    "Make a bold wrapped header and align quantity cells to the right; print the style properties.",
    r'''
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    sheet = Workbook().active
    sheet.append(["Long item description", "Quantity"])
    sheet.append(["Book", 3])
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].alignment = Alignment(wrap_text=True)
    sheet["B2"].alignment = Alignment(horizontal="right")
    print(sheet["A1"].font.bold, sheet["A1"].alignment.wrap_text)
    print(sheet["B2"].alignment.horizontal)
    ''',
    "Save and reload a number format, proving it persists while the numeric value is unchanged.",
    r'''
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    book.active["A1"] = 12.5
    book.active["A1"].number_format = "0.00"
    book.save("format.xlsx")
    restored = load_workbook("format.xlsx")
    print(restored.active["A1"].value, restored.active["A1"].number_format)
    restored.close()
    ''',
    "Readable formatting is separate from data correctness. Use explicit styles and widths; keep actual values and units clear. Reopen output to verify styles survive.",
    [
        ("Does number_format change the stored value?", "No; it changes display.", "Yes, always.", "Only with a fill.", "The numeric value and the display format are distinct properties."),
        ("How should a shared style be changed?", "Assign a new style object.", "Mutate arbitrary style fields in place.", "Convert the cell to a string.", "Style objects are designed for assignment/copy rather than in-place shared mutation."),
    ])

add(1, "2. Formulas and the cached-value trap",
    "Write Excel formulas while keeping a trustworthy Python-computed total.",
    [
        ("Formula string", "A cell string beginning = is a formula. It describes a calculation for Excel/compatible software, rather than being evaluated by openpyxl."),
        ("data_only=True", "On loading, data_only returns last stored cached formula results, if present. It does not recalculate formulas; a newly written formula may load as None."),
        ("Python total vs workbook formula", "A Python calculation can supply a verified summary immediately; a formula provides a viewer-recalculable worksheet expression. Keep those roles distinct."),
    ],
    "Write =SUM(B2:B3) for the spreadsheet viewer, but compute the same total in Python when creating a PDF or validating output. "
    "openpyxl does not evaluate the formula. A freshly generated workbook normally has no cached result, so data_only=True can return None. "
    "Do not mistake a stale cache for current truth. For our automated summaries, validated raw values and Python totals are authoritative.",
    'sheet["B4"] = "=SUM(B2:B3)"\npython_total = sum(validated_values)\n# data_only reads a stored cache; it does not run Excel',
    r'''
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    sheet = book.active
    sheet.append(["Item", "Quantity"])
    sheet.append(["Book", 3])
    sheet.append(["Pen", 5])
    sheet["B4"] = "=SUM(B2:B3)"
    book.save("formulas.xlsx")
    formulas = load_workbook("formulas.xlsx", data_only=False)
    cached = load_workbook("formulas.xlsx", data_only=True)
    print(formulas.active["B4"].value)
    print(cached.active["B4"].value)
    print("Python total:", 3 + 5)
    formulas.close()
    cached.close()
    ''',
    "Formula text सुरक्षित भयो, तर cached result None छ—यो code broken होइन, calculation engine चलेको छैन। "
    "Python total 8 भएकाले PDF summary बनाउन Excel खोल्नु आवश्यक छैन। "
    "Existing file को cache पनि पुरानो हुन सक्छ, त्यसैले important imports मा raw values/formula policy तय गर।",
    ("load_workbook(data_only=True) recalculates every formula",
     "It only reads cached results; compute totals independently when needed",
     "The library is a file editor, not Excel's calculation engine."),
    "Write a row formula quantity*unit_price and compute the matching Python value.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet.append(["Quantity", "Unit price", "Amount"])
    sheet.append([3, 20, "=A2*B2"])
    print(sheet["C2"].value)
    print(sheet["A2"].value * sheet["B2"].value)
    ''',
    "Reject formula cells in a strict raw-data import instead of silently treating their missing/stale cache as a trustworthy number.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    sheet["A1"] = "=1+2"
    sheet["A2"] = 3
    for row in sheet.iter_rows(max_col=1, max_row=2):
        cell = row[0]
        print(cell.coordinate, "reject formula" if cell.data_type == "f" else "raw value accepted")
    ''',
    "openpyxl writes formulas but never calculates them. data_only reads a cache, not a fresh result. Keep independently verified totals and define whether formula input is allowed.",
    [
        ("What may a new formula return with data_only=True?", "None because no cached calculation exists.", "Always the computed number.", "The formula string always.", "No calculation engine has populated the result cache."),
        ("What should drive a newly generated PDF total here?", "Python's total from validated raw rows", "An assumed formula cache", "Cell colour", "The verified source values can be computed without opening Excel."),
    ])

add(1, "3. Conditional formatting and useful controls",
    "Highlight meaningful thresholds and preserve filters/frozen headers.",
    [
        ("CellIsRule", "A conditional formatting rule tells the spreadsheet viewer to style cells when a condition holds. It is not input validation."),
        ("freeze_panes; auto_filter.ref", "freeze_panes keeps preceding rows/columns visible while scrolling. auto_filter.ref gives the viewer a filterable range; it does not filter rows in Python."),
    ],
    "Highlight quantities below 3 with a labelled policy and readable colour. Keep the header visible using A2: only row 1 stays frozen. "
    "AutoFilter enables controls in a viewer but does not remove rows from your dataset. "
    "Conditional colours are rendered by Excel/compatible software; our Python checks verify the stored rule, not a screenshot of Excel.",
    'sheet.conditional_formatting.add("B2:B4", CellIsRule(operator="lessThan", formula=["3"], fill=fill))\nsheet.freeze_panes = "A2"\nsheet.auto_filter.ref = "A1:B4"',
    r'''
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.formatting.rule import CellIsRule
    book = Workbook()
    sheet = book.active
    sheet.append(["Item", "Quantity"])
    for row in [("Book", 2), ("Pen", 5), ("Bag", 1)]:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="FFFFE2CC")
    sheet.conditional_formatting.add("B2:B4", CellIsRule(
        operator="lessThan", formula=["3"], fill=fill))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:B4"
    book.save("controls.xlsx")
    restored = load_workbook("controls.xlsx")
    print(restored.active.freeze_panes, restored.active.auto_filter.ref)
    print(len(restored.active.conditional_formatting))
    restored.close()
    ''',
    "Threshold 3 भन्दा कम भए highlight हुने instruction file मा छ। Freeze/filter पनि save भएर फर्किए। "
    "AutoFilter राख्यो भन्दैमा Python total बाट filtered rows स्वतः हट्दैनन्। Colour लाई explanation बिना meaning नदेऊ।",
    ("Use cell colour as the only validation or omit filtered rows automatically",
     "Validate values separately and define aggregation filters explicitly",
     "Viewer formatting/filter controls do not change your Python data pipeline."),
    "Compare A2 and B2 freeze settings; state that B2 also freezes the first column.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    for position in ("A2", "B2"):
        sheet.freeze_panes = position
        print(sheet.freeze_panes)
    print("A2: row 1; B2: row 1 and column A")
    ''',
    "Compute low-stock records in Python separately from the conditional formatting rule.",
    r'''
    rows = [("Book", 2), ("Pen", 5), ("Bag", 1)]
    threshold = 3
    low = [item for item, quantity in rows if quantity < threshold]
    print(low)
    ''',
    "Formatting highlights data; it does not validate or filter it. Save explicit viewer controls and compute report selections with clear Python rules.",
    [
        ("Does auto_filter.ref filter data during Python iteration?", "No", "Yes", "Only with colours", "It records a viewer control; your Python code still reads the underlying rows."),
        ("What does freeze_panes='A2' freeze?", "The first row", "The first two rows", "The whole sheet", "Cells above the designated position remain visible."),
    ])

add(1, "4. Print layout and a monthly report",
    "Design a compact report with clear units, printable bounds and a verified total.",
    [
        ("print_area; page_setup; print_title_rows", "These define printable cells, paper orientation/scaling and repeated header rows. They affect printing/export layout, not calculations."),
        ("Fit-to-width vs fit-to-height", "Fit to one page wide while allowing multiple pages high for long tables; forcing everything onto one page can make text unreadable."),
    ],
    "Monthly reports should identify the month, data units and total. Choose a landscape layout for a wider table, "
    "repeat the header on each printed page, and define a print area that excludes unrelated cells. "
    "Print settings are viewer instructions; openpyxl does not create a rendered PDF from an Excel workbook. "
    "Day 5 uses ReportLab to create a PDF directly from the same validated data.",
    'sheet.print_area = "A1:B4"\nsheet.print_title_rows = "1:1"\nsheet.page_setup.orientation = "landscape"\nsheet.page_setup.fitToWidth = 1\nsheet.page_setup.fitToHeight = 0',
    r'''
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "October"
    sheet.append(["Department", "Units"])
    sheet.append(["Books", 8])
    sheet.append(["Stationery", 5])
    sheet.append(["Total", 13])
    sheet.print_area = "A1:B4"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    print(sheet.title, sheet["B4"].value)
    print(sheet.page_setup.orientation, sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight)
    print(sheet.print_title_rows)
    ''',
    "Report title October हो र units explicitly लेखिएका छन्। Total actual numeric 13 हो। "
    "एक page wide तर unlimited height हुँदा long table readability बचाउन सकिन्छ। "
    "Spreadsheet PDF मा render गर्ने viewer र ReportLab बाट नयाँ PDF बनाउने काम अलग हुन्।",
    ("Force a long report to one tiny page and treat print settings as a PDF export",
     "Keep text readable and generate/render PDF through an appropriate tool",
     "Print metadata does not render a document, and excessive scaling can make a report unusable."),
    "Create a report with units for two departments and an explicit numeric total.",
    r'''
    from openpyxl import Workbook
    sheet = Workbook().active
    rows = [("Books", 8), ("Stationery", 5)]
    sheet.append(["Department", "Units"])
    for row in rows:
        sheet.append(row)
    sheet.append(["Total", sum(value for name, value in rows)])
    print(sheet["A4"].value, sheet["B4"].value)
    ''',
    "Save/reload landscape and repeated-header settings.",
    r'''
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    sheet = book.active
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    book.save("print-layout.xlsx")
    restored = load_workbook("print-layout.xlsx")
    print(restored.active.print_title_rows, restored.active.page_setup.orientation)
    restored.close()
    ''',
    "A useful report has labels, units, a verified total and readable print settings. Distinguish stored print metadata from actual document rendering.",
    [
        ("Does openpyxl's page_setup directly render a PDF?", "No", "Yes", "Only in landscape", "It stores printing instructions in the workbook."),
        ("Why allow more than one page high for a long table?", "To avoid shrinking text into unreadability.", "To change totals.", "To remove headers.", "Readability matters more than a forced single-page fit."),
    ])

add(2, "1. Native Excel charts: data and category references",
    "Create a column chart with correct values, labels and header boundaries.",
    [
        ("BarChart; type='col'", "A native Excel bar-chart object can use vertical columns when type is col. It is stored in the workbook rather than as a screenshot."),
        ("Reference; add_data; set_categories", "Reference selects worksheet ranges. add_data supplies numeric series; set_categories supplies matching labels."),
        ("titles_from_data=True", "When the data range includes the header, this uses its first cell as the series title instead of plotting it as a number."),
        ("zipfile.ZipFile for verification", "An .xlsx is a ZIP-based package. Inspecting its stored chart parts can verify chart creation without relying on private openpyxl object fields."),
    ],
    "Put category labels in column A and quantities in B. The data range B1:B3 includes its header; the categories A2:A3 contain two labels. "
    "Keep the category and numeric row counts aligned. Exclude the total row from categories, because it would double-count visually. "
    "Attach the chart away from the table so neither obscures the other. Viewer appearance can differ; verify data and inspect the report in a viewer.",
    'chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)\nchart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))\nsheet.add_chart(chart, "D2")',
    r'''
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from zipfile import ZipFile
    book = Workbook()
    sheet = book.active
    for row in [("Category", "Units"), ("Books", 8), ("Pens", 5)]:
        sheet.append(row)
    chart = BarChart()
    chart.type = "col"
    chart.title = "Units by category"
    chart.y_axis.title = "Units"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    book.save("chart.xlsx")
    with ZipFile("chart.xlsx") as archive:
        parts = [name for name in archive.namelist() if name.startswith("xl/charts/chart") and name.endswith(".xml")]
        print("stored charts:", len(parts))
    ''',
    "Header data range मा छ, categories मा छैन। दुई values का लागि दुई labels छन्। "
    "Chart D2 मा राखिएको छ। ZIP check ले chart stored भएको प्रमाण दिन्छ; यसले viewer को visual appearance सही छ भन्ने पूर्ण प्रमाण दिँदैन।",
    ("Include the grand total as another category", "Chart only the comparable component categories",
     "A total is an aggregate of those components; displaying it as another peer category can mislead."),
    "Construct matching value/category ranges for three categories and print the reference strings.",
    r'''
    from openpyxl import Workbook
    from openpyxl.chart import Reference
    sheet = Workbook().active
    for row in [("Category", "Units"), ("A", 2), ("B", 3), ("C", 4)]:
        sheet.append(row)
    print(Reference(sheet, min_col=2, min_row=1, max_row=4))
    print(Reference(sheet, min_col=1, min_row=2, max_row=4))
    ''',
    "Make a one-series column chart and verify its workbook contains a chart part after saving.",
    r'''
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from zipfile import ZipFile
    book = Workbook()
    sheet = book.active
    sheet.append(["Item", "Units"])
    sheet.append(["Book", 8])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=2))
    sheet.add_chart(chart, "D2")
    book.save("one-chart.xlsx")
    with ZipFile("one-chart.xlsx") as archive:
        print("xl/charts/chart1.xml" in archive.namelist())
    ''',
    "A chart is only as correct as its ranges. Include headers intentionally, align labels/values, exclude totals as peer categories and verify stored chart parts plus visual output.",
    [
        ("Where should category labels usually come from here?", "A2:A3", "B1:B3", "The grand total only", "Categories use the text labels without the header."),
        ("What does titles_from_data=True do?", "Uses the included header as a series title.", "Calculates totals.", "Adds one value automatically.", "It prevents a title cell being treated as a data point."),
    ])

add(2, "2. Grouped summaries in Python",
    "Compute category totals from raw rows before creating a chart.",
    [
        ("Grouped aggregation", "Collect records sharing a key and combine their numeric values. A dictionary maps each category to its running total."),
        ("Stable category order", "Sort category keys before writing rows so repeated runs use the same table/chart ordering."),
    ],
    "Raw rows may repeat Books several times. A chart of repeated raw labels is not the same as a category-total report. "
    "First validate records, then aggregate using totals.get(category,0)+units. Sort the categories, write a summary sheet, and chart that table. "
    "This is a pivot-style summary computed in Python, not a native interactive Excel PivotTable.",
    'totals[category] = totals.get(category, 0) + units\nfor category in sorted(totals):\n    summary.append([category, totals[category]])',
    r'''
    from openpyxl import Workbook
    rows = [("Books", 3), ("Pens", 5), ("Books", 4)]
    totals = {}
    for category, units in rows:
        totals[category] = totals.get(category, 0) + units
    sheet = Workbook().active
    sheet.append(["Category", "Units"])
    for category in sorted(totals):
        sheet.append([category, totals[category]])
    print(list(sheet.iter_rows(values_only=True)))
    print("grand total:", sum(totals.values()))
    ''',
    "Books को 3 र 4 जोडिएर 7 भयो; Pens 5 हो। Output sorted भएकाले repeated runs मा क्रम स्थिर हुन्छ। "
    "Raw total र grouped grand total 12 मिल्नुपर्छ। Native PivotTable click-to-reconfigure गर्न यो code ले सिर्जना गरेको होइन।",
    ("Claim a static grouped table is a native interactive PivotTable",
     "Describe it accurately as a Python-computed pivot-style summary",
     "A generated aggregation is not an Excel PivotTable object with its own cache and interactive controls."),
    "Group quantities by department and check that grouped totals sum to the raw total.",
    r'''
    rows = [("A", 2), ("B", 5), ("A", 3)]
    totals = {}
    for department, amount in rows:
        totals[department] = totals.get(department, 0) + amount
    assert sum(totals.values()) == sum(amount for department, amount in rows)
    print(sorted(totals.items()))
    ''',
    "Group by month in YYYY-MM format and output chronological lexical order.",
    r'''
    rows = [("2026-10", 4), ("2026-09", 3), ("2026-10", 2)]
    totals = {}
    for month, amount in rows:
        totals[month] = totals.get(month, 0) + amount
    print(sorted(totals.items()))
    ''',
    "Aggregate before charting when the question is grouped totals. Keep order stable and verify that group sums reconcile with the raw data.",
    [
        ("Is this generated summary a native interactive Excel PivotTable?", "No", "Yes", "Only if sorted", "It is an ordinary worksheet populated by Python aggregation."),
        ("What reconciliation check is useful?", "Grouped totals sum to the raw total.", "Every category has the same amount.", "The chart is colourful.", "The grouping should neither lose nor double-count records."),
    ])

add(2, "3. A two-dimensional pivot-style matrix",
    "Summarise month by category and define missing combinations explicitly.",
    [
        ("Composite key", "A tuple such as (month, category) identifies a two-dimensional group in a dictionary."),
        ("Matrix missing-cell policy", "For validated sales totals, no recorded sale for a month/category is displayed as zero here. Missing source measurements are not universally zero."),
    ],
    "Collect all months and categories, total each (month,category) combination, then write one matrix row per month. "
    "This report uses zero for absent recorded-sales combinations, a declared policy. If records mean unknown measurements, that policy would be wrong. "
    "Add row totals as a separate numeric column. Understand the grouping before trying to automate a viewer's native pivot interface.",
    'totals[(month, category)] = totals.get((month, category), 0) + units\nvalues = [totals.get((month, category), 0) for category in categories]',
    r'''
    from openpyxl import Workbook
    rows = [("2026-09", "Books", 3), ("2026-10", "Books", 4), ("2026-10", "Pens", 5)]
    months = sorted({month for month, category, units in rows})
    categories = sorted({category for month, category, units in rows})
    totals = {}
    for month, category, units in rows:
        key = (month, category)
        totals[key] = totals.get(key, 0) + units
    sheet = Workbook().active
    sheet.append(["Month"] + categories + ["Total"])
    for month in months:
        values = [totals.get((month, category), 0) for category in categories]
        sheet.append([month] + values + [sum(values)])
    for row in sheet.iter_rows(values_only=True):
        print(row)
    ''',
    "Tuple key ले month र category दुवै राख्छ। September/Pens record नभएकाले policy अनुसार 0 आउँछ। "
    "यो unknown value लाई 0 बनाएको होइन; validated sales logs मा absent combination को अर्थ तय गरिएको छ।",
    ("Treat every absent measurement as zero without a policy",
     "Declare whether absence means zero, unknown or invalid",
     "Zero is a real numeric value; missing data may mean something different."),
    "Aggregate the same category in two months using tuple keys.",
    r'''
    rows = [("Sep", "Books", 2), ("Oct", "Books", 3), ("Sep", "Books", 4)]
    totals = {}
    for month, category, amount in rows:
        key = (month, category)
        totals[key] = totals.get(key, 0) + amount
    print(sorted(totals.items()))
    ''',
    "Compute row totals in a small matrix and assert that the grand total is 12.",
    r'''
    matrix = [[3, 0], [4, 5]]
    row_totals = [sum(row) for row in matrix]
    assert sum(row_totals) == 12
    print(row_totals, sum(row_totals))
    ''',
    "Composite keys produce clear two-dimensional summaries. Declare missing-combination semantics and reconcile matrix/row/grand totals against the source.",
    [
        ("Which key identifies a month/category combination?", "(month, category)", "category alone", "The last row number", "A tuple retains both grouping dimensions."),
        ("Does an absent measurement always mean zero?", "No", "Yes", "Only in Excel", "Missing-data meaning belongs to the data contract."),
    ])

add(2, "4. Chart choice: categories versus a time sequence",
    "Use bar/column charts for categories and a labelled line chart for ordered months.",
    [
        ("LineChart", "A line chart connects ordered observations, such as consecutive months. It can imply sequence; do not use that implication carelessly for unordered categories."),
        ("Independent chart objects", "Create a separate chart object for each chart you attach. A chart's ranges and titles should describe that report's question."),
    ],
    "A column chart compares category totals. A line chart can show a monthly trend when the categories are in chronological order. "
    "Our month labels are YYYY-MM and sorted, so their order is meaningful. Axis labels must say the units. "
    "A chart cannot repair an incorrect total or missing-data policy. Keep a readable numeric table beside it.",
    'chart = LineChart()\nchart.title = "Monthly units"\nchart.y_axis.title = "Units"\nchart.add_data(values, titles_from_data=True)\nchart.set_categories(month_labels)',
    r'''
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from zipfile import ZipFile
    book = Workbook()
    sheet = book.active
    for row in [("Month", "Units"), ("2026-09", 3), ("2026-10", 9)]:
        sheet.append(row)
    chart = LineChart()
    chart.title = "Monthly units"
    chart.y_axis.title = "Units"
    chart.x_axis.title = "Month"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    book.save("trend.xlsx")
    with ZipFile("trend.xlsx") as archive:
        print("xl/charts/chart1.xml" in archive.namelist())
    print("monthly total:", 3 + 9)
    ''',
    "Month क्रम September अनि October छ। Line जोड्नु अर्थपूर्ण छ किनकि sequence छ। "
    "Books/Pens जस्तो unordered categories मा त्यसरी trend भन्नु misleading हुन्छ। Axis units र numeric table सँगै राख।",
    ("Connect arbitrary categories with a line and call it a time trend",
     "Choose a chart whose visual implication matches the data",
     "A connecting line suggests order/continuity that arbitrary categories may not have."),
    "Sort out-of-order months before constructing a trend table.",
    r'''
    monthly = {"2026-10": 9, "2026-08": 2, "2026-09": 3}
    print(sorted(monthly.items()))
    ''',
    "Create two independent chart objects for a category comparison and a month trend.",
    r'''
    from openpyxl.chart import BarChart, LineChart
    category_chart = BarChart()
    trend_chart = LineChart()
    print(type(category_chart).__name__, type(trend_chart).__name__)
    print(category_chart is trend_chart)
    ''',
    "Match the chart to the question, preserve numeric tables and label axes with units. Correct ranges, order and aggregation matter more than decoration.",
    [
        ("Which chart naturally shows an ordered monthly sequence?", "LineChart", "Any chart with random order", "A pie chart only", "Connecting ordered time observations communicates a trend."),
        ("Can a chart fix incorrect aggregation?", "No", "Yes", "Only with a title", "The visual uses the values you supply, so correctness must come first."),
    ])

PDF_FIXTURE = clean(r'''
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4
canvas = Canvas("source.pdf", pagesize=A4)
canvas.drawString(50, 780, "Page 1: Python study")
canvas.showPage()
canvas.drawString(50, 780, "Page 2: daily practice")
canvas.save()
''')

add(3, "1. Read PDF pages and extract available text",
    "Inspect a generated text PDF and understand when extraction cannot supply text.",
    [
        ("pypdf.PdfReader; pages", "PdfReader reads a PDF structure. pages is an indexed collection; the first page is pages[0], unlike 1-based spreadsheet coordinates."),
        ("extract_text; OCR limitation", "extract_text retrieves text represented in the PDF; pypdf is not OCR. An image-only scan may yield empty/no text even when a human can read the image."),
        ("ReportLab Canvas: fixture setup", "Canvas creates a PDF. drawString places a text line at an (x,y) position; showPage completes one page, and save finishes the file. This small recipe provides our test input; Day 5 develops layout in detail."),
    ],
    "Each snippet creates its own tiny source PDF, so there is no missing sample download. The six setup lines place two plain English text lines on two pages. "
    "Then use PdfReader and extract_text. PDF text order may not match a paragraph's visual order; a PDF is not a Word document or table database. "
    "Scans may need OCR, which is a separate tool/workflow. Do not equate an empty extraction result with an empty-looking page.",
    'reader = PdfReader("source.pdf")\nfor page_number, page in enumerate(reader.pages, 1):\n    text = page.extract_text() or ""',
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader
    reader = PdfReader("source.pdf")
    print("pages:", len(reader.pages))
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        print(number, text.strip())
    '''),
    "Canvas setup ले source बनायो; PdfReader ले त्यही file पढ्यो। enumerate(...,1) display page number दिन्छ, "
    "तर indexed access अझै zero-based हो। extract_text or '' ले absent text लाई stringमा राख्छ; "
    "यो OCR चलाएको होइन। Complex layout मा extracted reading order जाँच्नुपर्छ।",
    ("pypdf always reads text from scanned page images",
     "Use a separate OCR workflow when the PDF has no extractable text",
     "Text extraction reads PDF text objects, not the words inside raster images."),
    "Extract the first page only and print its text and word count under a simple whitespace-token policy.",
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader
    text = PdfReader("source.pdf").pages[0].extract_text() or ""
    print(text.strip())
    print("whitespace tokens:", len(text.split()))
    '''),
    "Create a blank PDF page and demonstrate that extractable text can be empty.",
    r'''
    from pypdf import PdfWriter, PdfReader
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write("blank.pdf")
    reader = PdfReader("blank.pdf")
    print(repr(reader.pages[0].extract_text() or ""))
    ''',
    "PDF page indexes start at 0. Text extraction is neither OCR nor guaranteed visual reading order. Handle empty results and verify the document's actual content/layout.",
    [
        ("What is the index of the first PDF page?", "0", "1", "A1", "The pages collection follows ordinary Python indexing."),
        ("Does pypdf perform OCR on scanned images?", "No", "Yes", "Only with extract_text", "OCR is a separate process; image-only pages may not expose text."),
    ])

add(3, "2. Split and merge with explicit page ranges",
    "Select pages without off-by-one errors and write to separate files.",
    [
        ("PdfWriter; append; add_page", "PdfWriter creates output PDFs. append adds a document/page selection; add_page adds a selected PageObject."),
        ("Half-open page range", "append(..., pages=(start, stop)) uses zero-based, stop-exclusive bounds; (0,1) selects only the first page."),
        ("Page-order policy", "A merge preserves the order you request. A filename sort is not automatically the intended document order."),
    ],
    "Split pages into new files and merge deliberate sequences, leaving source files untouched. "
    "For our two-page input, (0,1) contains page 1 and (1,2) contains page 2. "
    "Plain-document merging does not guarantee preservation of every interactive form, signature or navigation feature. "
    "For real forms, consider field-name collisions and verification separately.",
    'writer = PdfWriter()\nwriter.append("source.pdf", pages=(0, 1))\nwriter.write("first-page.pdf")',
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader, PdfWriter
    first = PdfWriter()
    first.append("source.pdf", pages=(0, 1))
    first.write("first.pdf")
    second = PdfWriter()
    second.append("source.pdf", pages=(1, 2))
    second.write("second.pdf")
    merged = PdfWriter()
    merged.append("first.pdf")
    merged.append("second.pdf")
    merged.write("merged.pdf")
    result = PdfReader("merged.pdf")
    print("merged pages:", len(result.pages))
    print([page.extract_text().strip() for page in result.pages])
    '''),
    "Source का दुई pages छुट्टाछुट्टै outputs भए। Append क्रम first अनि second भएकाले merged order उस्तै रह्यो। "
    "Page count मात्रै जाँचेर content/order नमिलेको छुट्न सक्छ; extracted labels पनि तुलना गर।",
    ("pages=(1,2) selects the first and second pages", "It selects only zero-based index 1",
     "The stop bound is excluded; page selection follows a half-open interval."),
    "Reverse a two-page document into a new PDF and verify its text order.",
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader("source.pdf")
    writer = PdfWriter()
    for page in reversed(reader.pages):
        writer.add_page(page)
    writer.write("reversed.pdf")
    print([page.extract_text().strip() for page in PdfReader("reversed.pdf").pages])
    '''),
    "Merge three independently generated one-page PDFs in a declared list order and verify three pages.",
    r'''
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader, PdfWriter
    paths = []
    for number in range(1, 4):
        path = f"part-{number}.pdf"
        canvas = Canvas(path)
        canvas.drawString(50, 750, f"Part {number}")
        canvas.save()
        paths.append(path)
    writer = PdfWriter()
    for path in paths:
        writer.append(path)
    writer.write("three-parts.pdf")
    print([page.extract_text().strip() for page in PdfReader("three-parts.pdf").pages])
    ''',
    "Specify page ranges and merge order explicitly, use new files, and verify counts plus content/order. Interactive/signature preservation is a separate contract.",
    [
        ("How many pages does pages=(0,1) select?", "One", "Two", "None", "The stop index is excluded."),
        ("What determines append merge order?", "The sequence of your append calls", "Always alphabetical filenames", "Page colours", "The writer follows the order you supply."),
    ])

add(3, "3. Rotation and metadata are separate from text",
    "Rotate a selected page and inspect descriptive document metadata.",
    [
        ("PageObject.rotate", "rotate changes page rotation in multiples of 90 degrees. A viewing transform is not rewriting the text into a new paragraph layout."),
        ("metadata; add_metadata", "Metadata may include title/author and can be absent. PdfWriter.add_metadata writes keys such as /Title; metadata is not proof of who really created the file."),
        ("mediabox", "The media box describes the page's physical bounds in points. A rotation does not automatically mean those stored bounds swap."),
    ],
    "Use a selected page, apply a 90-degree rotation and save a new output. "
    "Check the resulting rotation property and page count; inspect visually for the intended orientation. "
    "Metadata describes a document but can be modified. Do not treat an author string as authentication. "
    "Later overlays/merges with rotated pages need careful coordinate handling; this part performs only a rotation.",
    'page = reader.pages[0]\npage.rotate(90)\nwriter.add_page(page)\nwriter.add_metadata({"/Title": "Practice"})',
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader("source.pdf")
    page = reader.pages[0]
    page.rotate(90)
    writer = PdfWriter()
    writer.add_page(page)
    writer.add_metadata({"/Title": "Rotated practice"})
    writer.write("rotated.pdf")
    restored = PdfReader("rotated.pdf")
    print(len(restored.pages), restored.pages[0].rotation)
    print(restored.metadata.get("/Title"))
    print("text present:", "Python" in (restored.pages[0].extract_text() or ""))
    '''),
    "Rotation 90 stored छ, तर text अझै Python नै हो। Title metadata पनि read-back भयो। "
    "यही property देखियो भन्दैमा visual placement सबै ठीक छ भन्ने होइन; actual page हेर्नुपर्छ।",
    ("A metadata author field proves the author's identity", "Treat metadata as editable descriptive information",
     "Metadata is not a signature or authentication mechanism."),
    "Rotate the first page 180 degrees and verify that rotation value after saving.",
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader, PdfWriter
    page = PdfReader("source.pdf").pages[0]
    page.rotate(180)
    writer = PdfWriter()
    writer.add_page(page)
    writer.write("upside-down.pdf")
    print(PdfReader("upside-down.pdf").pages[0].rotation)
    '''),
    "Write and inspect /Title and /Author metadata on a blank-page practice PDF.",
    r'''
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": "Study", "/Author": "Demo"})
    writer.write("metadata.pdf")
    metadata = PdfReader("metadata.pdf").metadata
    print(metadata.get("/Title"), metadata.get("/Author"))
    ''',
    "Rotation is a viewing transform; metadata is descriptive and editable. Verify saved properties and inspect orientation visually before trusting the output.",
    [
        ("What rotation increments are supported here?", "Multiples of 90 degrees", "Any arbitrary fraction", "Only zero", "Page rotation operates in quarter turns."),
        ("Does metadata authenticate an author?", "No", "Yes", "Only with /Title", "It is editable descriptive information."),
    ])

add(3, "4. Encrypted, malformed and empty-text documents",
    "Handle expected failures honestly and avoid inventing extracted content.",
    [
        ("is_encrypted; decrypt", "An encrypted PDF may need an authorised password before reading. decrypt reports whether access succeeded; do not assume a call succeeded."),
        ("PdfReadError", "pypdf's PDF-read exception covers expected malformed/empty-file reading failures. Catch intended errors and keep source files for diagnosis."),
        ("Extraction completeness", "A page count or some extracted text is not proof that every visual word/table was recovered. Preserve page boundaries and note extraction limits."),
    ],
    "Use only the toy password in this isolated example; it is not a request for your real credentials. "
    "Check decrypt's result before accessing pages, and do not log real passwords. Some encryption algorithms need additional cryptographic dependencies; "
    "the simple demo uses pypdf's default supported mode. A decrypted reader may still report is_encrypted=True because the original file is encrypted. "
    "For malformed documents, report a read failure rather than saving an invented empty output.",
    'if reader.is_encrypted:\n    if reader.decrypt(authorised_password) == 0:\n        raise ValueError("password rejected")',
    r'''
    from pypdf import PdfWriter, PdfReader
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("practice-only")
    writer.write("locked.pdf")
    reader = PdfReader("locked.pdf")
    print(reader.is_encrypted)
    print(reader.decrypt("wrong-demo") == 0)
    print(reader.decrypt("practice-only") != 0)
    print("accessible pages:", len(reader.pages))
    ''',
    "पहिलो True file encrypted भएको बताउँछ। गलत demo password असफल हुन्छ; सही demo password सफल भएपछि page access हुन्छ। "
    "यो authorization bypass होइन। आफ्नो documents मा approved access मात्र प्रयोग गर र real password examples/sourceमा नराख।",
    ("Assume decryption worked or invent text for an unreadable file",
     "Check access/read results and report specific failure",
     "A failed access or extraction is not permission to fabricate document content or erase the original."),
    "Catch reading an empty byte stream as a PDF and print a clear failure.",
    r'''
    from io import BytesIO
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        PdfReader(BytesIO(b""))
    except PdfReadError:
        print("empty/malformed PDF rejected")
    ''',
    "Extract each page separately, retaining page numbers, then calculate a simple total whitespace-token count.",
    PDF_FIXTURE + clean(r'''
    from pypdf import PdfReader
    counts = []
    for number, page in enumerate(PdfReader("source.pdf").pages, 1):
        count = len((page.extract_text() or "").split())
        counts.append(count)
        print(number, count)
    print("total whitespace tokens:", sum(counts))
    '''),
    "Check authorised decryption, catch expected read errors and preserve originals. Text extraction may be incomplete; keep page boundaries and never fabricate missing content.",
    [
        ("What does decrypt returning 0 mean?", "Password/access was not accepted.", "The PDF has zero pages.", "The file was deleted.", "Do not read encrypted pages as though decryption succeeded."),
        ("Does some extracted text prove complete visual recovery?", "No", "Yes", "Only with metadata", "PDF layout and image-based content can limit extraction."),
    ])

add(4, "1. ReportLab canvas: coordinates, pages and finishing",
    "Create a simple PDF with an explicit page size and coordinate system.",
    [
        ("PDF point; A4; reportlab.lib.units.mm", "PDF positions use points, 72 per inch. A4 supplies width/height in points; mm converts millimetres to that unit."),
        ("Bottom-left origin", "Default canvas coordinates start at the bottom left. Bigger y moves upwards; page_height - top_margin gives a top-based placement."),
        ("setFont; drawString; line; showPage; save", "Canvas drawing methods place content manually. showPage completes a page and resets drawing state; save finalises the document."),
    ],
    "Day 4 used a tiny Canvas recipe to create test inputs. Now learn its geometry rather than guessing y coordinates. "
    "A4 is portrait width/height; place a heading 20 mm from its top. Set the font on each page because showPage resets state. "
    "Canvas is useful for fixed layouts, but manually placing long paragraphs/tables requires wrapping and pagination logic. "
    "The next part uses Platypus for flowing document content.",
    'width, height = A4\ncanvas.drawString(20*mm, height-20*mm, "Heading")\ncanvas.showPage()\ncanvas.save()',
    r'''
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from pypdf import PdfReader
    width, height = A4
    canvas = Canvas("layout.pdf", pagesize=A4)
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(20*mm, height-20*mm, "Study report")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(20*mm, height-32*mm, "Units studied: 12")
    canvas.line(20*mm, height-38*mm, width-20*mm, height-38*mm)
    canvas.save()
    reader = PdfReader("layout.pdf")
    print("pages:", len(reader.pages))
    print("Study report" in (reader.pages[0].extract_text() or ""))
    print("A4 points:", round(width), round(height))
    ''',
    "Top margin बाट राख्न height - margin प्रयोग भयो। y ठूलो हुँदा माथि जान्छ, तल होइन। "
    "save पछि file पूरा हुन्छ। Extracted text/checks useful छन्, तर काटिएको वा overlap भएको छैन कि actual PDF render गरेर हेर्नुपर्छ।",
    ("Treat y as distance down from the top by default", "Default canvas y grows upwards from the bottom",
     "Incorrect coordinate assumptions can place content off-page or overlap it."),
    "Create two pages and reset a readable font after each showPage.",
    r'''
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader
    canvas = Canvas("two-pages.pdf")
    for number in (1, 2):
        canvas.setFont("Helvetica", 12)
        canvas.drawString(50, 750, f"Page {number}")
        canvas.showPage()
    canvas.save()
    print([page.extract_text().strip() for page in PdfReader("two-pages.pdf").pages])
    ''',
    "Convert a 20 mm margin to points and place text inside an A4 page with it.",
    r'''
    from reportlab.lib.units import mm
    from reportlab.lib.pagesizes import A4
    margin = 20 * mm
    print(round(margin, 2))
    print(0 < margin < A4[0] / 2)
    print("top y:", round(A4[1] - margin, 2))
    ''',
    "Use explicit page size, units and margins. Canvas defaults to a bottom-left origin; showPage resets drawing state. Finish with save and inspect the rendered layout.",
    [
        ("Where is the default canvas origin?", "Bottom left", "Top left", "Page centre", "Increasing y moves upwards."),
        ("Which method finalises the PDF file?", "save", "drawString", "setFont", "Drawing commands alone do not finish the document."),
    ])

add(4, "2. Platypus: paragraphs, tables and safe text",
    "Use flowing layout for a readable summary instead of manual line coordinates.",
    [
        ("SimpleDocTemplate; story; Flowable", "A document template lays out a sequence of Flowable objects. Paragraph, Spacer and Table are common pieces of that story."),
        ("Paragraph markup; escape", "ReportLab Paragraph understands a small markup syntax. Escape arbitrary text before embedding it so data such as & or < is not interpreted as markup."),
        ("Table; TableStyle; repeatRows", "Table arranges rows/columns; TableStyle specifies presentation. repeatRows=1 repeats the first header row if the table splits across pages."),
    ],
    "Use paragraph styles for a title and body, then a table with deliberate widths. Table cells should use Paragraph for long text that must wrap. "
    "Escaping text is especially important when item names contain &, < or >. A basic string in a table does not automatically solve long-word layout. "
    "Set margins and font sizes, repeat headers and inspect actual pages. Tables need enough room; keep a numeric total outside or clearly separated from detail rows.",
    'story = [Paragraph(title, styles["Title"]), Spacer(1,12), Table(rows, repeatRows=1)]\nSimpleDocTemplate(path, pagesize=A4).build(story)',
    r'''
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from xml.sax.saxutils import escape
    from pypdf import PdfReader
    styles = getSampleStyleSheet()
    rows = [["Item", "Units"], [Paragraph(escape("Books & pens"), styles["BodyText"]), "12"]]
    table = Table(rows, colWidths=[300, 100], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story = [Paragraph("Study summary", styles["Title"]), Spacer(1, 12), table]
    SimpleDocTemplate("summary.pdf", pagesize=A4).build(story)
    text = PdfReader("summary.pdf").pages[0].extract_text() or ""
    print("Study summary" in text, "Books & pens" in text)
    ''',
    "Story ले title, gap र table क्रमैसँग राख्छ। TableStyle tuples मा coordinates (column,row) हुन्; -1 अन्तिम column/row हो। "
    "Paragraph लाई escaped data दिएकाले ampersand content भएर सुरक्षित देखिन्छ। Widths जोड्दा available page width ननाघ्नुपर्छ।",
    ("Put arbitrary user text directly into Paragraph markup", "Escape data text before building Paragraphs",
     "Markup-sensitive characters can change interpretation or cause parse/layout errors."),
    "Build a paragraph containing literal 'A < B & C' and verify its extracted text.",
    r'''
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from xml.sax.saxutils import escape
    from pypdf import PdfReader
    text = "A < B & C"
    SimpleDocTemplate("escaped.pdf").build([Paragraph(escape(text), getSampleStyleSheet()["BodyText"])])
    print(text in (PdfReader("escaped.pdf").pages[0].extract_text() or ""))
    ''',
    "Create a two-row table with a bold header and verify both labels survive PDF generation.",
    r'''
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from pypdf import PdfReader
    table = Table([["Item", "Units"], ["Book", 3]], colWidths=[250, 100], repeatRows=1)
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    SimpleDocTemplate("table.pdf").build([table])
    text = PdfReader("table.pdf").pages[0].extract_text() or ""
    print("Item" in text, "Book" in text)
    ''',
    "Flowing layout handles wrapping/pagination more naturally than manual coordinates. Escape data used in markup, choose explicit widths, repeat headers and inspect the final pages.",
    [
        ("What is a story in Platypus?", "An ordered list of layout Flowables", "The PDF's password", "A raw CSV file", "The template places its Paragraph/Table/Spacer objects in sequence."),
        ("Why escape arbitrary Paragraph text?", "It uses markup-sensitive syntax.", "To encrypt it.", "To remove all spaces.", "Escaping preserves literal data characters."),
    ])

add(4, "3. An invoice from integer minor units",
    "Compute line amounts once and display a consistent total in the PDF.",
    [
        ("Integer minor-unit amounts", "For this toy invoice, amounts are stored in hundredths of a unit. Integer arithmetic avoids binary-float rounding in the line/total calculation."),
        ("Shared calculation model", "One validated data model should drive the detail table and total; independently recomputing different report totals invites inconsistencies."),
    ],
    "The example is arithmetic and document layout, not tax/accounting guidance. Three books cost 1250 minor units each and two pens cost 250 each. "
    "Compute line totals by quantity*unit_minor, sum once and format with whole/remainder divided by 100. "
    "Reject negative amounts or invalid quantities before generating real output. Avoid invented tax rates or assuming a currency has the same minor-unit policy everywhere.",
    'line_minor = quantity * unit_minor\ntotal_minor = sum(line_amounts)\nlabel = f"{minor // 100}.{minor % 100:02d}"',
    r'''
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table
    from reportlab.lib.styles import getSampleStyleSheet
    from pypdf import PdfReader
    items = [("Book", 3, 1250), ("Pen", 2, 250)]
    def money(minor):
        return f"{minor // 100}.{minor % 100:02d}"
    total = sum(quantity * price for name, quantity, price in items)
    rows = [["Item", "Qty", "Unit", "Line"]]
    for name, quantity, price in items:
        rows.append([name, str(quantity), money(price), money(quantity * price)])
    rows.append(["Total (demo units)", "", "", money(total)])
    styles = getSampleStyleSheet()
    SimpleDocTemplate("invoice.pdf").build([
        Paragraph("Practice invoice", styles["Title"]),
        Table(rows, colWidths=[180, 50, 90, 90], repeatRows=1),
    ])
    print("minor total:", total)
    print("display total:", money(total))
    print(money(total) in (PdfReader("invoice.pdf").pages[0].extract_text() or ""))
    ''',
    "3×1250 + 2×250 = 4250 minor units हो, display 42.50। सबै calculations integerमा भए। "
    "PDF total त्यही shared value बाट आयो; अलग float formula चलाएको होइन। Real business rules आफ्नो context अनुसार छुट्टै तय हुन्छन्।",
    ("Round each inconsistent float calculation and hope totals agree",
     "Use a declared amount model and calculate shared totals once",
     "Representation and rounding policies must be explicit; formatting alone is not arithmetic correctness."),
    "Compute the total of two items using integer minor units, including a zero-cost item.",
    r'''
    items = [("Notebook", 2, 1500), ("Gift", 1, 0)]
    total = sum(quantity * unit_minor for name, quantity, unit_minor in items)
    print(total)
    print(f"{total // 100}.{total % 100:02d}")
    ''',
    "Validate quantity and unit_minor as actual nonnegative integers, rejecting booleans and negative values.",
    r'''
    def checked(quantity, unit_minor):
        if type(quantity) is not int or type(unit_minor) is not int or quantity < 0 or unit_minor < 0:
            raise ValueError("nonnegative integer amounts required")
        return quantity * unit_minor
    for pair in ((2, 1500), (True, 100), (2, -1)):
        try:
            print(checked(*pair))
        except ValueError:
            print("rejected")
    ''',
    "Separate input validation, amount representation, arithmetic and display. Use one set of verified totals throughout the report; document its units and assumptions.",
    [
        ("What is the worked invoice total in minor units?", "4250", "4000", "42", "Three 1250-unit books plus two 250-unit pens total 4250."),
        ("Does a 0.00 display format establish a money arithmetic policy?", "No", "Yes", "Only in PDF", "Representation, rounding and calculation rules must be chosen separately."),
    ])

add(4, "4. Long reports, page numbers and font coverage",
    "Paginate a long table and understand that Unicode text needs suitable fonts and shaping.",
    [
        ("repeatRows and page callbacks", "A long Table can split across pages with repeated headers. onFirstPage/onLaterPages callbacks draw consistent page-number footers."),
        ("pdfmetrics.registerFont; TTFont", "Register a compatible TrueType font to embed its supported glyphs. A font file's availability/licence and glyph coverage need checking."),
        ("Unicode glyphs vs shaping", "Storing Unicode does not guarantee readable glyphs or correct script shaping. Nepali/complex scripts and Japanese may need different fonts and a verified shaping/rendering workflow."),
    ],
    "Build a long table with a repeatable header and enough margins for a footer. Use page callbacks rather than hard-coded page counts. "
    "ReportLab's bundled Vera font provides a reproducible embedded-font example for text such as café; it is not presented as a complete Nepali/Japanese font. "
    "Our main report prototype explicitly uses English/ASCII labels; it rejects unsupported PDF text instead of silently drawing broken glyphs. "
    "For Nepali/Japanese output, choose suitable licensed fonts, verify coverage/shaping and visually inspect every generated page.",
    'document.build(story, onFirstPage=footer, onLaterPages=footer)\npdfmetrics.registerFont(TTFont("StudyFont", font_path))',
    r'''
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.pagesizes import A4
    from pypdf import PdfReader
    rows = [["Record", "Units"]] + [[f"Item {i}", str(i)] for i in range(1, 81)]
    table = Table(rows, colWidths=[300, 100], repeatRows=1)
    table.setStyle(TableStyle([("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.drawString(50, 25, f"Page {document.page}")
        canvas.restoreState()
    SimpleDocTemplate("long-report.pdf", pagesize=A4,
                     topMargin=50, bottomMargin=50).build(
        [table], onFirstPage=footer, onLaterPages=footer)
    pages = PdfReader("long-report.pdf").pages
    print("multiple pages:", len(pages) > 1)
    print("header on all pages:", all("Record" in (page.extract_text() or "") for page in pages))
    print("last record present:", "Item 80" in (pages[-1].extract_text() or ""))
    ''',
    "Long table स्वतः pages मा विभाजित हुन्छ; repeatRows ले header दोहोर्‍याउँछ। Footer मा document.page actual number हो। "
    "saveState/restoreState ले callback का font changes body layoutमा असर नपार्ने बनाउँछन्। "
    "Text extraction tests पछि actual pages पनि inspect गर; glyph/overlap issues text checks ले छुटाउन सक्छन्।",
    ("Unicode input guarantees any PDF font will render Nepali/Japanese correctly",
     "Verify glyph coverage, script shaping and rendered output",
     "A basic Latin font cannot supply every script, and glyph coverage alone may not establish correct complex-script shaping."),
    "Embed ReportLab's bundled Vera font and write café; verify extracted Unicode text.",
    r'''
    import reportlab
    from pathlib import Path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader
    font_path = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    pdfmetrics.registerFont(TTFont("StudyVera", str(font_path)))
    canvas = Canvas("font-demo.pdf")
    canvas.setFont("StudyVera", 12)
    canvas.drawString(50, 750, "café")
    canvas.save()
    print("café" in (PdfReader("font-demo.pdf").pages[0].extract_text() or ""))
    ''',
    "Check the declared ASCII-only prototype policy for Book and सीता rather than silently corrupting unsupported PDF text.",
    r'''
    def supported_prototype_label(value):
        return isinstance(value, str) and value.isascii() and all(32 <= ord(ch) <= 126 for ch in value)
    for value in ("Book", "सीता"):
        print(value, supported_prototype_label(value))
    print("Non-ASCII output needs a deliberate font/shaping extension.")
    ''',
    "Long reports need repeated headers, footers and visual checks. Font embedding supports only its covered glyphs; complex-script output needs a tested shaping/font workflow, not an unsupported promise.",
    [
        ("How can a long table repeat its header?", "repeatRows=1", "One giant font", "data_only=True", "The table repeats its first row on split pages."),
        ("Does Vera demonstrate complete Nepali/Japanese support?", "No", "Yes", "Only with save", "This example covers supported Latin text; other scripts require suitable fonts and verification."),
    ])

add(5, "1. Word documents: paragraphs, runs, styles and tables",
    "Create and read a modest .docx document while separating content from formatting.",
    [
        ("python-docx; Document", "Install the package named python-docx, then import Document from docx. It reads/writes .docx, not old binary .doc files; no Word installation is needed for file creation."),
        ("Paragraph; Run", "A paragraph groups text; a run is a text segment with its own formatting. A single visible sentence can be split into several runs."),
        ("Heading styles; add_table", "Semantic heading styles support document structure. add_table creates rows/cells; assigning cell.text replaces that cell's content."),
    ],
    "Use heading styles for a title and sections rather than turning every line into a large bold run. "
    "A paragraph can contain plain and bold runs. Tables keep related data in cells; .docx is not plain text where every field can be replaced globally. "
    "Unicode text can be stored, but readable rendering still depends on appropriate fonts and the document viewer. "
    "Use a new output filename and inspect the generated document before sharing it.",
    'document = Document()\ndocument.add_heading("Study report", level=1)\nparagraph = document.add_paragraph("Name: ")\nparagraph.add_run("Sita").bold = True\ndocument.save("report.docx")',
    r'''
    from docx import Document
    document = Document()
    document.add_heading("Study report", level=1)
    paragraph = document.add_paragraph("Learner: ")
    paragraph.add_run("Sita").bold = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Subject"
    table.rows[0].cells[1].text = "Units"
    cells = table.add_row().cells
    cells[0].text = "Python"
    cells[1].text = "12"
    document.save("study.docx")
    restored = Document("study.docx")
    print([p.text for p in restored.paragraphs])
    print([[cell.text for cell in row.cells] for row in restored.tables[0].rows])
    ''',
    "Learner paragraph दुई runsबाट बनेको छ: plain label र bold name। Read-back paragraph.text ले संयुक्त text दिन्छ। "
    "Table text paragraph listमा स्वतः आउँदैन; tables अलग iterate गर्नुपर्छ। Save भयो भन्दैमा सबै layout ठीक भयो भन्ने होइन।",
    ("document.paragraphs automatically includes every table/header text",
     "Read paragraphs, tables and other supported document regions explicitly",
     "Different document regions have separate APIs; a simple paragraph loop is not a complete text inventory."),
    "Create a paragraph with two runs and print the combined text and run-level bold flags.",
    r'''
    from docx import Document
    paragraph = Document().add_paragraph()
    paragraph.add_run("Hello ")
    paragraph.add_run("Sita").bold = True
    print(paragraph.text)
    print([run.bold for run in paragraph.runs])
    ''',
    "Save and restore a simple Nepali paragraph, then explain that stored Unicode does not verify viewer font shaping.",
    r'''
    from docx import Document
    document = Document()
    document.add_paragraph("सीता")
    document.save("unicode.docx")
    restored = Document("unicode.docx")
    print(restored.paragraphs[0].text)
    print(restored.paragraphs[0].text == "सीता")
    ''',
    "Use structural styles, understand run-level formatting and read tables separately. File creation/text preservation and visual document quality are different checks.",
    [
        ("Which package is installed for from docx import Document?", "python-docx", "A package necessarily named docx", "pypdf", "The distribution is python-docx and the import namespace is docx."),
        ("Can one visible sentence contain several runs?", "Yes", "No", "Only in tables", "Run boundaries reflect formatting and editing, not necessarily word boundaries."),
    ])

add(5, "2. Templates and ten personalised letters from CSV",
    "Use a deliberately simple template and handle split placeholders without claiming rich-format preservation.",
    [
        ("Template loading", "Document(template_path) opens a .docx template; load a fresh copy for each recipient so one letter does not inherit another's values."),
        ("Split placeholder problem", "A placeholder such as {{name}} may be split across several runs. Replacing individual run strings can miss it."),
        ("Whole-paragraph replacement trade-off", "Replacing paragraph.text can replace a split placeholder but resets run-level formatting. Our supported template uses plain body paragraphs only, not arbitrary rich documents."),
        ("Safe filename policy", "Use checked numeric/string IDs for output names, not arbitrary recipient text that could contain path separators."),
    ],
    "Generate a fresh document per record. The simple template contains a plain body paragraph with {{name}} and {{id}}. "
    "Use paragraph.text replacement for this supported template and state the formatting loss; a rich template needs a run-aware templating strategy. "
    "This exercise does not replace headers, footers, text boxes or fields. Validate CSV names/IDs and avoid overwriting existing letters. "
    "Build an isolated folder for the ten-letter exercise; do not send the letters anywhere.",
    'for record in records:\n    document = Document(template_path)\n    for paragraph in document.paragraphs:\n        paragraph.text = paragraph.text.replace("{{name}}", record["name"])\n    document.save(new_path)',
    r'''
    import csv
    import io
    from pathlib import Path
    from docx import Document
    template = Document()
    paragraph = template.add_paragraph()
    paragraph.add_run("Dear {{na")
    paragraph.add_run("me}}, your learner ID is {{id}}.")
    template.save("template.docx")
    records = list(csv.DictReader(io.StringIO("id,name\n007,Sita\n008,Hari\n")))
    outputs = []
    for record in records:
        if not record["id"].isascii() or not record["id"].isdigit() or not record["name"].strip():
            raise ValueError("invalid recipient")
        document = Document("template.docx")
        for paragraph in document.paragraphs:
            paragraph.text = paragraph.text.replace("{{name}}", record["name"]).replace("{{id}}", record["id"])
        path = Path("letter-" + record["id"] + ".docx")
        if path.exists():
            raise ValueError("output exists")
        document.save(path)
        outputs.append(Document(path).paragraphs[0].text)
    print(outputs)
    ''',
    "Template मा name placeholder दुई runs मा फुटेको छ। Paragraph text बाट replace गर्दा भेटियो, तर run formatting reset हुन्छ। "
    "Fresh Document प्रत्येक loop मा load भएकोले Sita/Hari मिसिँदैनन्। Output IDs check गरिएका छन्। "
    "यो simple-body template मात्रै हो; सबै Word formatsको universal mail merge होइन।",
    ("Replace each run and promise all split placeholders/rich formatting are preserved",
     "Define a simple template contract or use a tested run-aware templating system",
     "Visible placeholders can cross run boundaries; paragraph.text replacement trades run formatting for simple replacement."),
    "Generate ten letters from self-contained CSV data and assert every saved document contains its own recipient ID.",
    r'''
    import csv
    import io
    from pathlib import Path
    from docx import Document
    csv_text = "id,name\n" + "".join(f"{number:03d},Learner {number}\n" for number in range(1, 11))
    template = Document()
    template.add_paragraph("Dear {{name}}, ID {{id}}.")
    template.save("template.docx")
    paths = []
    for record in csv.DictReader(io.StringIO(csv_text)):
        document = Document("template.docx")
        document.paragraphs[0].text = document.paragraphs[0].text.replace(
            "{{name}}", record["name"]).replace("{{id}}", record["id"])
        path = Path(f"letter-{record['id']}.docx")
        document.save(path)
        assert record["id"] in Document(path).paragraphs[0].text
        paths.append(path)
    print("verified letters:", len(paths))
    ''',
    "Show why run-level replacement misses a split placeholder while paragraph-level replacement finds it.",
    r'''
    from docx import Document
    paragraph = Document().add_paragraph()
    paragraph.add_run("{{na")
    paragraph.add_run("me}}")
    for run in paragraph.runs:
        run.text = run.text.replace("{{name}}", "Sita")
    print(paragraph.text)
    paragraph.text = paragraph.text.replace("{{name}}", "Sita")
    print(paragraph.text)
    ''',
    "Load fresh templates, validate recipient IDs/names and verify every output. The simple paragraph approach handles split placeholders by sacrificing run formatting; richer templates require a separate tested strategy.",
    [
        ("Why load the template anew for each recipient?", "To avoid carrying previous replacements into the next letter.", "To send mail automatically.", "To change CSV delimiters.", "Each document starts from the unreplaced template."),
        ("What can paragraph.text replacement reset?", "Run-level formatting", "The Python interpreter", "All filesystem names", "The paragraph is reconstructed as simple text, not all its original formatted runs."),
    ])

PROJECT = clean(r'''
import argparse
import csv
from datetime import date
from pathlib import Path
import re
import sys
import tempfile
from zipfile import BadZipFile, ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from xml.sax.saxutils import escape
from pypdf import PdfReader

FIELDS = ["date", "department", "item", "units"]

def validate_rows(rows):
    if not rows or len(rows) > 10000:
        raise ValueError("require 1..10000 records")
    checked = []
    for number, row in enumerate(rows, 1):
        if not isinstance(row, dict) or set(row) != set(FIELDS):
            raise ValueError(f"record {number}: fields must be date/department/item/units")
        text = row["date"]
        if not isinstance(text, str) or re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", text) is None:
            raise ValueError(f"record {number}: ISO date required")
        date.fromisoformat(text)
        labels = {}
        for field in ("department", "item"):
            value = row[field]
            if not isinstance(value, str) or not value.strip() or len(value) > 60:
                raise ValueError(f"record {number}: invalid {field}")
            value = value.strip()
            if not value.isascii() or any(not 32 <= ord(ch) <= 126 for ch in value):
                raise ValueError(f"record {number}: prototype PDF requires printable ASCII labels")
            if value.startswith(("=", "+", "-", "@")):
                raise ValueError(f"record {number}: formula-like label not supported")
            labels[field] = value
        units = row["units"]
        if type(units) is not int or not 0 <= units <= 100000:
            raise ValueError(f"record {number}: units must be integer 0..100000")
        checked.append({"date": text, **labels, "units": units})
    return checked

def read_rows(source, kind):
    source = Path(source)
    if source.stat().st_size > 10_000_000:
        raise ValueError("input exceeds local 10 MB policy")
    rows = []
    if kind == "csv":
        with source.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, strict=True)
            if reader.fieldnames != FIELDS:
                raise ValueError("wrong CSV header")
            for number, row in enumerate(reader, 1):
                if None in row or any(value is None for value in row.values()):
                    raise ValueError(f"record {number}: wrong column count")
                if re.fullmatch(r"[0-9]{1,6}", row["units"].strip()) is None:
                    raise ValueError(f"record {number}: integer units text required")
                row["units"] = int(row["units"].strip())
                rows.append(row)
                if len(rows) > 10000:
                    raise ValueError("too many records")
    elif kind == "xlsx":
        book = load_workbook(source, read_only=True, data_only=False)
        try:
            if "Sales" not in book.sheetnames:
                raise ValueError("Sales worksheet required")
            sheet = book["Sales"]
            if sheet.max_column != 4 or sheet.max_row > 10001:
                raise ValueError("expected a small four-column Sales table")
            header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=4))
            if [cell.value for cell in header] != FIELDS:
                raise ValueError("wrong Excel header")
            for cells in sheet.iter_rows(min_row=2, max_col=4):
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("formula inputs not supported")
                rows.append(dict(zip(FIELDS, [cell.value for cell in cells])))
        finally:
            book.close()
    else:
        raise ValueError("input format must be csv or xlsx")
    return validate_rows(rows)

def summarise(rows):
    totals = {}
    for row in rows:
        key = row["department"]
        totals[key] = totals.get(key, 0) + row["units"]
    return {"records": len(rows), "units": sum(row["units"] for row in rows),
            "departments": sorted(totals.items())}

def write_excel(path, rows, summary):
    book = Workbook()
    raw = book.active
    raw.title = "Raw"
    raw.append(FIELDS)
    for row in rows:
        raw.append([row[field] for field in FIELDS])
    grouped = book.create_sheet("Summary")
    grouped.append(["Department", "Units"])
    for department, units in summary["departments"]:
        grouped.append([department, units])
    metrics = book.create_sheet("Metrics")
    metrics.append(["Metric", "Value"])
    metrics.append(["Records", summary["records"]])
    metrics.append(["Total units", summary["units"]])
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF16324F")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in ("A", "B", "C", "D"):
            sheet.column_dimensions[column].width = 24
    chart = BarChart()
    chart.type = "col"
    chart.title = "Units by department"
    chart.y_axis.title = "Units"
    last = len(summary["departments"]) + 1
    chart.add_data(Reference(grouped, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(grouped, min_col=1, min_row=2, max_row=last))
    grouped.add_chart(chart, "D2")
    book.save(path)

def write_pdf(path, summary):
    styles = getSampleStyleSheet()
    story = [Paragraph("Office automation report", styles["Title"]),
             Paragraph(f"Records: {summary['records']} | Total units: {summary['units']}",
                       styles["BodyText"]), Spacer(1, 16)]
    data = [["Department", "Units"]]
    for department, units in summary["departments"]:
        data.append([Paragraph(escape(department), styles["BodyText"]), str(units)])
    table = Table(data, colWidths=[330, 100], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 1), (1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(table)
    story.extend([Spacer(1, 14), Paragraph(
        "Generated from validated source rows. Units are counts, not money.",
        styles["BodyText"])])
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.drawString(50, 25, f"Page {document.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=50, rightMargin=50,
        topMargin=50, bottomMargin=50).build(
            story, onFirstPage=footer, onLaterPages=footer)

def build_report(source, kind, destination):
    destination = Path(destination)
    if destination.exists():
        raise ValueError("output folder exists; choose a new folder")
    rows = read_rows(source, kind)
    summary = summarise(rows)
    # Stage both files; publish their directory only after validation succeeds.
    with tempfile.TemporaryDirectory(prefix=".office-stage-", dir=destination.parent) as folder:
        stage = Path(folder)
        write_excel(stage / "report.xlsx", rows, summary)
        write_pdf(stage / "summary.pdf", summary)
        book = load_workbook(stage / "report.xlsx", read_only=True)
        try:
            if book["Metrics"]["B3"].value != summary["units"]:
                raise ValueError("Excel total verification failed")
        finally:
            book.close()
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(stage / "summary.pdf").pages)
        if f"Total units: {summary['units']}" not in pdf_text:
            raise ValueError("PDF total verification failed")
        stage.replace(destination)
    return summary

def main(argv=None):
    parser = argparse.ArgumentParser(description="Generate a validated Excel and PDF report")
    parser.add_argument("source")
    parser.add_argument("output_folder")
    parser.add_argument("--input-format", required=True, choices=["csv", "xlsx"])
    args = parser.parse_args(argv)
    try:
        summary = build_report(args.source, args.input_format, args.output_folder)
    except (OSError, ValueError, csv.Error, BadZipFile, InvalidFileException, SyntaxError) as error:
        print("Report failed:", error, file=sys.stderr)
        return 1
    print(f"Report ready: {summary['records']} records, {summary['units']} units")
    return 0
''')

SAMPLE = clean(r'''
records = [
    {"date": "2026-10-01", "department": "Books", "item": "Python book", "units": 3},
    {"date": "2026-10-02", "department": "Stationery", "item": "Pen", "units": 5},
    {"date": "2026-10-03", "department": "Books", "item": "Notebook", "units": 4},
]
''')

CREATE_INPUTS = clean(r'''
with tempfile.TemporaryDirectory() as folder:
    folder = Path(folder)
    source_csv = folder / "sales.csv"
    with source_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)
    source_xlsx = folder / "sales.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(FIELDS)
    for row in records:
        sheet.append([row[field] for field in FIELDS])
    book.save(source_xlsx)
    for source, kind in ((source_csv, "csv"), (source_xlsx, "xlsx")):
        summary = build_report(source, kind, folder / ("result-" + kind))
        print(kind, summary)
''')

PIPELINE_FAILURES = clean(r'''
with tempfile.TemporaryDirectory() as folder:
    folder = Path(folder)
    source = folder / "bad.csv"
    source.write_text("date,department,item,units\n2026-02-31,Books,Book,3\n", encoding="utf-8")
    output = folder / "output"
    try:
        build_report(source, "csv", output)
    except ValueError:
        print("bad date rejected")
    assert not output.exists()
    output.mkdir()
    marker = output / "keep.txt"
    marker.write_text("KEEP", encoding="utf-8")
    try:
        build_report(source, "csv", output)
    except ValueError:
        print("existing output folder protected")
    assert marker.read_text(encoding="utf-8") == "KEEP"
''')

PIPELINE_CHECKS = clean(r'''
checked = validate_rows(records)
summary = summarise(checked)
assert summary["records"] == 3
assert summary["units"] == 12
assert summary["departments"] == [("Books", 7), ("Stationery", 5)]
assert sum(units for department, units in summary["departments"]) == summary["units"]
for bad in (True, -1, 2.5):
    changed = [{**records[0], "units": bad}]
    try:
        validate_rows(changed)
    except ValueError:
        continue
    raise AssertionError("invalid units accepted")
print("counts, group totals, reconciliation and unit rejections passed")
''')

add(5, "3. Complete project: one CSV/Excel-to-report pipeline",
    "Validate a source table once and produce matching Excel and PDF reports in a new folder.",
    [
        ("Canonical model and adapters", "CSV/Excel readers convert a declared source schema into the same list of validated dictionaries. Both report writers use one shared summary."),
        ("Staging directory", "Generate and verify both files in a temporary sibling directory, then rename it to the new output folder. A failed generation should not publish half a report."),
        ("Formula-like input policy", "This prototype rejects labels beginning =/+/-/@ and Excel formula cells. Spreadsheet imports should have an explicit policy for text/formulas rather than accidentally executing expressions."),
        ("Small local-file contract", "The project supports local CSV UTF-8 or a small four-column Sales worksheet. It is not arbitrary workbook preservation, a hostile-upload service or concurrent-writer coordination."),
    ],
    "Input columns are date,department,item,units. Dates use YYYY-MM-DD and must exist; units are actual integers 0..100000. "
    "Printable ASCII department/item labels are 1..60 characters after trimming because the prototype PDF uses basic Latin fonts. "
    "Non-ASCII labels fail clearly; extend the font/shaping workflow from Day 5 before promising Nepali/Japanese PDF output. "
    "CSV uses an optional UTF-8 BOM; XLSX requires Sales and raw, non-formula values. Blank Excel rows are skipped only if all four cells are empty. "
    "A local input cap of 10 MB and 10000 records limits this teaching workload; file size alone is not a ZIP-bomb defence. "
    "Outputs are report.xlsx with Raw/Summary/Metrics/chart and summary.pdf with the same totals. Existing output folders are refused. "
    "This staged directory strategy assumes a single writer and does not promise filesystem transactions across concurrent programs.",
    'python office_report.py sales.csv output-october --input-format csv\npython office_report.py sales.xlsx output-excel --input-format xlsx\npython office_report.py --help',
    PROJECT + SAMPLE + CREATE_INPUTS,
    "पूरा project पाँच layersमा पढ: read_rows → validate_rows → summarise → write_excel/write_pdf → build_report/main। "
    "CSV र Excel फरक readers भए पनि canonical rows एउटै छन्। Summary एक पटक निकालेर दुवै report मा प्रयोग भयो। "
    "Excel formula cache प्रयोग गरिएको छैन; raw inputs बाट Python ले total 12 निकाल्छ। "
    "Stagingमा दुवै files लेखेर read-back total जाँचिन्छ; सफल भएपछि नयाँ output folder publish हुन्छ। "
    "Source file overwrite हुँदैन। Complete runnable script docs/python-week24-office-report.py र sample CSV पनि repoमा छन्।",
    ("Write Excel first into final output, then leave it behind if PDF generation fails",
     "Stage both files, verify them, then publish the new report folder",
     "A partial pair is misleading. Staging lets a failure clean up temporary output without damaging an existing report."),
    "Prove an impossible date publishes no output, and an existing output folder retains its marker unchanged.",
    PROJECT + PIPELINE_FAILURES,
    "Assert record count, total units, group totals and rejection of bool/negative/fractional units.",
    PROJECT + SAMPLE + PIPELINE_CHECKS,
    "One validated model and one shared summary keep outputs consistent. Protect sources, stage both reports and test failure-state preservation. Respect the declared input, font and concurrency limits.",
    [
        ("What supplies totals for both final reports?", "One Python summary from validated source rows", "A newly recalculated openpyxl cache", "PDF text guessed independently", "Shared calculation prevents inconsistent outputs."),
        ("What happens if the final output folder already exists?", "The pipeline refuses it without overwriting it.", "It is deleted automatically.", "Only the PDF is replaced.", "Choose a new output folder; existing reports stay protected."),
    ])

add(5, "4. Whole-week review: inspect files, not just success messages",
    "Rebuild the pipeline in stages and test both document content and visual layout.",
    [
        ("Structural vs visual verification", "Reopening files checks pages, values, text and chart parts. Rendering/opening in a viewer checks clipping, wrapping, glyphs and actual presentation. Neither replaces the other."),
        ("Generation boundary tests", "Test rejected inputs, existing outputs, read/write failure and source preservation, as well as a successful report pair."),
    ],
    "Take as many sessions as you need; Day 7 remains your existing review/rest day. "
    "Rebuild a workbook from five rows, add a formatted summary and chart, split a two-page PDF, create a flowing report and generate ten simple letters. "
    "Then run the full pipeline from CSV and XLSX and compare the same totals. "
    "Try a missing header, an impossible date, a boolean/negative quantity, a formula cell and an existing output folder. "
    "For PDF/Word/Excel presentation, open the actual generated document; a count of pages or a passing text assertion alone cannot establish visual quality.",
    'success checks: source → validated rows → shared totals → reopened output values\nfailure checks: intended error + no partial output + unchanged source/previous reports',
    r'''
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    book.active.append(["Item", "Units"])
    book.active.append(["Book", 3])
    book.save("review.xlsx")
    restored = load_workbook("review.xlsx")
    assert restored.active["B2"].value == 3
    restored.close()
    print("saved value verification passed")
    print("A visual viewer check is a separate step.")
    ''',
    "Read-back checks actualvalue verify गर्छन्। Viewer/renderमा column width, wrapping, font र clipping हेर्नुपर्छ। "
    "Long content राखेर पनि test गर; दुई छोटा rowsमा राम्रो देखिने layout ठूलो tableमा बिग्रिन सक्छ।",
    ("A file exists, therefore its content and layout are correct",
     "Reopen/check content and inspect actual rendering",
     "File creation alone does not establish values, page order, chart correctness or visual readability."),
    "Use the project to build a report from three CSV rows, then verify the Metrics total, stored chart and PDF text.",
    PROJECT + SAMPLE + clean(r'''
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        source = folder / "sales.csv"
        with source.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(records)
        output = folder / "report"
        build_report(source, "csv", output)
        book = load_workbook(output / "report.xlsx", read_only=True)
        assert book["Metrics"]["B3"].value == 12
        book.close()
        with ZipFile(output / "report.xlsx") as archive:
            assert "xl/charts/chart1.xml" in archive.namelist()
        text = "\n".join(page.extract_text() or "" for page in PdfReader(output / "summary.pdf").pages)
        assert "Total units: 12" in text
        print("Excel total, chart part and PDF total passed")
    '''),
    "Call the project's CLI entry point with an explicit argument list and verify its 0 success status.",
    PROJECT + SAMPLE + clean(r'''
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        source = folder / "sales.csv"
        with source.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(records)
        status = main([str(source), str(folder / "report"), "--input-format", "csv"])
        assert status == 0
        print("CLI status:", status)
    '''),
    "Week 24 mastery: read/validate Excel, style and chart correct totals, understand formula caches, extract/manipulate PDFs honestly, build readable PDF layouts, generate supported Word templates, and verify an end-to-end report pipeline. Do not rush weak parts.",
    [
        ("Does openpyxl calculate formulas?", "No", "Yes", "Only in read-only mode", "It stores formulas and reads caches; it is not Excel's calculation engine."),
        ("What are spreadsheet cell row indexes based on?", "1", "0", "-1", "A1 is row 1, column 1."),
        ("What are pypdf page indexes based on?", "0", "1", "A1", "The pages collection follows ordinary Python indexing."),
        ("Does values_only=True validate business records?", "No", "Yes", "Only with headers", "It returns values; schema/domain checks remain separate."),
        ("Is a Python group-by sheet a native interactive PivotTable?", "No", "Yes", "Only if it has a chart", "The generated table is a static Python-computed summary."),
        ("Does pypdf extract words from image-only scans by OCR?", "No", "Yes", "Only if metadata exists", "OCR is a separate workflow."),
        ("Where does the default ReportLab canvas origin start?", "Bottom left", "Top left", "Centre", "Coordinates grow upward from the default origin."),
        ("Why escape arbitrary Paragraph content?", "To preserve literal text instead of interpreting markup.", "To encrypt it.", "To calculate totals.", "Paragraph uses markup-sensitive characters."),
        ("What can simple paragraph.text template replacement lose?", "Run formatting", "All file extensions", "Python classes", "It replaces the paragraph's run structure with simpler text."),
        ("Does a passing text assertion prove a PDF's visual layout?", "No", "Yes", "Only if the title is present", "Rendering can still reveal clipping, overlap or missing glyphs."),
    ])

days["24.0"]["parts"][0]["sections"].insert(3, section("sh",
    "python -m pip install openpyxl pypdf reportlab python-docx",
    note="One-time setup in your practice Python environment. Package installation initially needs internet; all lesson examples use local files afterward. The Study Hub page itself remains offline."))
days["24.5"]["parts"][2]["sections"].append(section("sh",
    "python office_report.py sales.csv output-october --input-format csv\npython office_report.py --help",
    note="Save the complete project as office_report.py and use its sample sales.csv in your practice folder. The output folder must not already exist."))
days["24.5"]["parts"][-1]["sections"].append(section("p",
    "Optional official references: https://openpyxl.readthedocs.io/en/stable/tutorial.html · https://openpyxl.readthedocs.io/en/stable/charts/bar.html · https://pypdf.readthedocs.io/en/stable/user/extract-text.html · https://pypdf.readthedocs.io/en/stable/user/merging-pdfs.html · https://docs.reportlab.com/ · https://python-docx.readthedocs.io/en/latest/user/quickstart.html. The complete explanations and solutions are included here for offline study."))


def write():
    target = ROOT / "study-hub.html"
    original = target.read_bytes().decode("utf-8")
    start = original.index("const DAY_TEACH = {")
    end = original.index("\n};\n\nconst REST_DAY", start)
    if any(('"' + key + '":') in original[start:end] for key in days):
        raise RuntimeError("Week 24 exists; refusing to overwrite.")
    addition = "\n".join(json.dumps(k) + ": " + json.dumps(v, ensure_ascii=False, indent=2) + "," for k, v in days.items())
    prefix = original[:end].rstrip()
    if not prefix.endswith(","):
        prefix += ","
    result = prefix + "\n" + addition + original[end:]
    temporary = target.with_suffix(".html.tmp")
    temporary.write_bytes(result.encode("utf-8"))
    temporary.replace(target)
    (ROOT / "docs" / "python-week24-office-report.py").write_text(
        '"""Week 24 report pipeline. Python 3.11+.\n'
        'Install: python -m pip install openpyxl pypdf reportlab python-docx\n'
        'Input: local CSV UTF-8 or small Sales worksheet; exact four-column schema.\n'
        'Prototype PDF uses printable ASCII labels. Use a new output folder.\n'
        '"""\n' + PROJECT + '\nif __name__ == "__main__":\n    sys.exit(main())\n',
        encoding="utf-8")
    (ROOT / "docs" / "python-week24-sales.csv").write_text(
        "date,department,item,units\n2026-10-01,Books,Python book,3\n"
        "2026-10-02,Stationery,Pen,5\n2026-10-03,Books,Notebook,4\n",
        encoding="utf-8")
    parts = [p for day in days.values() for p in day["parts"]]
    quizzes = sum(len(s["lesson"]["quiz"]) for p in parts for s in p["sections"] if s["t"] == "checkpoint")
    print(f"Added Week 24: 6 days, {len(parts)} parts, {len(parts)*2} practice prompts, {quizzes} questions, {executed} verified examples.")


if __name__ == "__main__":
    write()

