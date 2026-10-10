"""Exercise actual report CLI output, failure-state preservation and staging."""
import csv
import importlib.util
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
sys.dont_write_bytecode = True

ROOT = Path(__file__).resolve().parent.parent
PROJECT = ROOT / "docs" / "python-week24-office-report.py"
passed = 0
preview = Path(tempfile.mkdtemp(prefix="studyhub-week24-preview-"))


def run(source, kind, destination, expected=0):
    global passed
    result = subprocess.run(
        [sys.executable, "-X", "utf8", str(PROJECT), str(source), str(destination),
         "--input-format", kind], capture_output=True, encoding="utf-8", timeout=15)
    assert result.returncode == expected, result.stderr
    if expected:
        assert result.stderr.startswith("Report failed:")
    else:
        assert "3 records, 12 units" in result.stdout
    passed += 1


def check(destination):
    book = load_workbook(destination / "report.xlsx")
    try:
        assert book.sheetnames == ["Raw", "Summary", "Metrics"]
        assert book["Metrics"]["B2"].value == 3
        assert book["Metrics"]["B3"].value == 12
        assert list(book["Summary"].iter_rows(min_row=2, values_only=True)) == [
            ("Books", 7), ("Stationery", 5)]
        assert book["Raw"].freeze_panes == "A2"
        assert book["Summary"]["A1"].font.bold
    finally:
        book.close()
    with ZipFile(destination / "report.xlsx") as archive:
        assert "xl/charts/chart1.xml" in archive.namelist()
    reader = PdfReader(destination / "summary.pdf")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert len(reader.pages) == 1
    assert "Total units: 12" in text and "Records: 3" in text
    assert "Books" in text and "Stationery" in text


with tempfile.TemporaryDirectory(prefix="studyhub-week24-tests-") as folder:
    folder = Path(folder)
    source = folder / "sales.csv"
    source.write_bytes((ROOT / "docs" / "python-week24-sales.csv").read_bytes())
    original = source.read_bytes()
    run(source, "csv", folder / "csv-output")
    check(folder / "csv-output")
    assert source.read_bytes() == original
    shutil.copy2(folder / "csv-output" / "summary.pdf", preview / "summary.pdf")
    shutil.copy2(folder / "csv-output" / "report.xlsx", preview / "report.xlsx")
    rows = list(csv.DictReader(source.open(encoding="utf-8", newline="")))
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["date", "department", "item", "units"])
    for row in rows:
        sheet.append([row["date"], row["department"], row["item"], int(row["units"])])
    xlsx = folder / "sales.xlsx"
    book.save(xlsx)
    xlsx_original = xlsx.read_bytes()
    run(xlsx, "xlsx", folder / "xlsx-output")
    check(folder / "xlsx-output")
    assert xlsx.read_bytes() == xlsx_original
    failures = [
        b"\xff",
        b"date,department,item,units\n",
        b"date,department,item,units\n2026-02-31,Books,Book,3\n",
        b"date,department,item,units\n2026-10-01,Books,Book,-1\n",
        b"date,department,item,units\n2026-10-01,Books,Book,2.5\n",
        b"date,department,item,units\n2026-10-01,Books,Book,100001\n",
        b"date,department,item,units\n2026-10-01,Books,Book\n",
        b"date,department,item,units\n2026-10-01,Books,Book,3,extra\n",
        b"date,department,item,units\n2026-10-01,=FORMULA,Book,3\n",
        "date,department,item,units\n2026-10-01,Books,सीता,3\n".encode("utf-8"),
        b"wrong,department,item,units\n",
    ]
    for index, content in enumerate(failures):
        bad = folder / f"bad-{index}.csv"
        output = folder / f"bad-output-{index}"
        bad.write_bytes(content)
        run(bad, "csv", output, expected=1)
        assert not output.exists()
        assert bad.read_bytes() == content
    for index, value in enumerate((True, -1, 2.5, "=1+2")):
        bad_book = Workbook()
        sheet = bad_book.active
        sheet.title = "Sales"
        sheet.append(["date", "department", "item", "units"])
        sheet.append(["2026-10-01", "Books", "Book", value])
        path = folder / f"bad-{index}.xlsx"
        bad_book.save(path)
        output = folder / f"bad-xlsx-output-{index}"
        run(path, "xlsx", output, expected=1)
        assert not output.exists()
    malformed = folder / "not-a-workbook.xlsx"
    malformed.write_bytes(b"not a ZIP file")
    run(malformed, "xlsx", folder / "malformed-output", expected=1)
    assert not (folder / "malformed-output").exists()
    existing = folder / "existing"
    existing.mkdir()
    (existing / "keep.txt").write_text("KEEP", encoding="utf-8")
    run(source, "csv", existing, expected=1)
    assert (existing / "keep.txt").read_text(encoding="utf-8") == "KEEP"
    spec = importlib.util.spec_from_file_location("week24_report", PROJECT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    def fail_pdf(path, summary):
        raise OSError("simulated PDF write failure")
    module.write_pdf = fail_pdf
    try:
        module.build_report(source, "csv", folder / "partial-output")
    except OSError:
        passed += 1
    else:
        raise AssertionError("simulated failure not propagated")
    assert not (folder / "partial-output").exists()
    assert not list(folder.glob(".office-stage-*"))
    assert source.read_bytes() == original

print(f"PASS: {passed} pipeline success/failure checks; CSV/XLSX totals, chart, PDF text, preserved sources/output and stage cleanup.")
print("Visual preview folder:", preview)
