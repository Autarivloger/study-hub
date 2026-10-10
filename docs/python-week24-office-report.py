"""Week 24 report pipeline. Python 3.11+.
Install: python -m pip install openpyxl pypdf reportlab python-docx
Input: local CSV UTF-8 or small Sales worksheet; exact four-column schema.
Prototype PDF uses printable ASCII labels. Use a new output folder.
"""
import argparse
import csv
from datetime import date
from pathlib import Path
import re
import sys
import tempfile
from zipfile import BadZipFile, ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from xml.sax.saxutils import escape
from pypdf import PdfReader

FIELDS = ["date", "department", "item", "units"]

def validate_rows(rows):
    if not rows or len(rows) > 10000:
        raise ValueError("require 1..10000 records")
    checked = []
    for number, row in enumerate(rows, 1):
        if not isinstance(row, dict) or set(row) != set(FIELDS):
            raise ValueError(f"record {number}: fields must be date/department/item/units")
        text = row["date"]
        if not isinstance(text, str) or re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", text) is None:
            raise ValueError(f"record {number}: ISO date required")
        date.fromisoformat(text)
        labels = {}
        for field in ("department", "item"):
            value = row[field]
            if not isinstance(value, str) or not value.strip() or len(value) > 60:
                raise ValueError(f"record {number}: invalid {field}")
            value = value.strip()
            if not value.isascii() or any(not 32 <= ord(ch) <= 126 for ch in value):
                raise ValueError(f"record {number}: prototype PDF requires printable ASCII labels")
            if value.startswith(("=", "+", "-", "@")):
                raise ValueError(f"record {number}: formula-like label not supported")
            labels[field] = value
        units = row["units"]
        if type(units) is not int or not 0 <= units <= 100000:
            raise ValueError(f"record {number}: units must be integer 0..100000")
        checked.append({"date": text, **labels, "units": units})
    return checked

def read_rows(source, kind):
    source = Path(source)
    if source.stat().st_size > 10_000_000:
        raise ValueError("input exceeds local 10 MB policy")
    rows = []
    if kind == "csv":
        with source.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, strict=True)
            if reader.fieldnames != FIELDS:
                raise ValueError("wrong CSV header")
            for number, row in enumerate(reader, 1):
                if None in row or any(value is None for value in row.values()):
                    raise ValueError(f"record {number}: wrong column count")
                if re.fullmatch(r"[0-9]{1,6}", row["units"].strip()) is None:
                    raise ValueError(f"record {number}: integer units text required")
                row["units"] = int(row["units"].strip())
                rows.append(row)
                if len(rows) > 10000:
                    raise ValueError("too many records")
    elif kind == "xlsx":
        book = load_workbook(source, read_only=True, data_only=False)
        try:
            if "Sales" not in book.sheetnames:
                raise ValueError("Sales worksheet required")
            sheet = book["Sales"]
            if sheet.max_column != 4 or sheet.max_row > 10001:
                raise ValueError("expected a small four-column Sales table")
            header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=4))
            if [cell.value for cell in header] != FIELDS:
                raise ValueError("wrong Excel header")
            for cells in sheet.iter_rows(min_row=2, max_col=4):
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("formula inputs not supported")
                rows.append(dict(zip(FIELDS, [cell.value for cell in cells])))
        finally:
            book.close()
    else:
        raise ValueError("input format must be csv or xlsx")
    return validate_rows(rows)

def summarise(rows):
    totals = {}
    for row in rows:
        key = row["department"]
        totals[key] = totals.get(key, 0) + row["units"]
    return {"records": len(rows), "units": sum(row["units"] for row in rows),
            "departments": sorted(totals.items())}

def write_excel(path, rows, summary):
    book = Workbook()
    raw = book.active
    raw.title = "Raw"
    raw.append(FIELDS)
    for row in rows:
        raw.append([row[field] for field in FIELDS])
    grouped = book.create_sheet("Summary")
    grouped.append(["Department", "Units"])
    for department, units in summary["departments"]:
        grouped.append([department, units])
    metrics = book.create_sheet("Metrics")
    metrics.append(["Metric", "Value"])
    metrics.append(["Records", summary["records"]])
    metrics.append(["Total units", summary["units"]])
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF16324F")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in ("A", "B", "C", "D"):
            sheet.column_dimensions[column].width = 24
    chart = BarChart()
    chart.type = "col"
    chart.title = "Units by department"
    chart.y_axis.title = "Units"
    last = len(summary["departments"]) + 1
    chart.add_data(Reference(grouped, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(grouped, min_col=1, min_row=2, max_row=last))
    grouped.add_chart(chart, "D2")
    book.save(path)

def write_pdf(path, summary):
    styles = getSampleStyleSheet()
    story = [Paragraph("Office automation report", styles["Title"]),
             Paragraph(f"Records: {summary['records']} | Total units: {summary['units']}",
                       styles["BodyText"]), Spacer(1, 16)]
    data = [["Department", "Units"]]
    for department, units in summary["departments"]:
        data.append([Paragraph(escape(department), styles["BodyText"]), str(units)])
    table = Table(data, colWidths=[330, 100], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 1), (1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(table)
    story.extend([Spacer(1, 14), Paragraph(
        "Generated from validated source rows. Units are counts, not money.",
        styles["BodyText"])])
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.drawString(50, 25, f"Page {document.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=50, rightMargin=50,
        topMargin=50, bottomMargin=50).build(
            story, onFirstPage=footer, onLaterPages=footer)

def build_report(source, kind, destination):
    destination = Path(destination)
    if destination.exists():
        raise ValueError("output folder exists; choose a new folder")
    rows = read_rows(source, kind)
    summary = summarise(rows)
    # Stage both files; publish their directory only after validation succeeds.
    with tempfile.TemporaryDirectory(prefix=".office-stage-", dir=destination.parent) as folder:
        stage = Path(folder)
        write_excel(stage / "report.xlsx", rows, summary)
        write_pdf(stage / "summary.pdf", summary)
        book = load_workbook(stage / "report.xlsx", read_only=True)
        try:
            if book["Metrics"]["B3"].value != summary["units"]:
                raise ValueError("Excel total verification failed")
        finally:
            book.close()
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(stage / "summary.pdf").pages)
        if f"Total units: {summary['units']}" not in pdf_text:
            raise ValueError("PDF total verification failed")
        stage.replace(destination)
    return summary

def main(argv=None):
    parser = argparse.ArgumentParser(description="Generate a validated Excel and PDF report")
    parser.add_argument("source")
    parser.add_argument("output_folder")
    parser.add_argument("--input-format", required=True, choices=["csv", "xlsx"])
    args = parser.parse_args(argv)
    try:
        summary = build_report(args.source, args.input_format, args.output_folder)
    except (OSError, ValueError, csv.Error, BadZipFile, InvalidFileException, SyntaxError) as error:
        print("Report failed:", error, file=sys.stderr)
        return 1
    print(f"Report ready: {summary['records']} records, {summary['units']} units")
    return 0

if __name__ == "__main__":
    sys.exit(main())
